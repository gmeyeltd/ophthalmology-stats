"""
ophthalmology-stats: compare prediction models evaluated on eyes, when some patients contribute both eyes, with
the patient as the unit of analysis.

Input: one row per eye, with a patient ID column and one column per model holding prediction errors (continuous
outputs; or predictions plus a reference column given as truth) or predicted probabilities plus the 0/1 diagnosis
(classifiers). Every pair of models is compared on each metric with a patient-level test, 95% confidence
interval and adjusted p-value. See README.md.

    import pandas as pd, ophthalmology_stats as osx
    res = osx.analyze(pd.read_csv("errors.csv"), subgroup="eye")
    osx.write_results_to_excel(res, "results.xlsx")

    python ophthalmology_stats.py errors.csv --subgroup eye --out results.xlsx

https://github.com/gmeyeltd/ophthalmology-stats (MIT license)
"""
__version__ = "1.3.0"

import re
import warnings
from itertools import combinations

import numpy as np
import pandas as pd
import scipy.stats as stats
from statsmodels.stats.multitest import multipletests

BOOTSTRAP_ITERATIONS = 4999       # p-values fall on a grid of 1/5000
BOOTSTRAP_SEED = 42               # fixed, so results are reproducible
CI_LEVEL = 0.95
MCNEMAR_THRESHOLD = 0.50          # default tolerance for the proportion of eyes with |error| <= within
MULTIPLE_TEST_METHOD = "romano_wolf"
FALLBACK_TEST_METHOD = "hommel"   # under Romano-Wolf, for tests without bootstrap null distributions
MIN_GROUP_SIZE = 5                # groups with fewer valid eyes are not tested
MEDIAN_WINDOW = 5                 # bootstrap medians: search window in sqrt(n) eyes (speed only; results are exact)

CONTINUOUS_METRICS = ["bias", "mae", "rmse", "medae", "sd", "within"]
BINARY_METRICS = ["auc", "brier", "logloss", "sensitivity", "specificity", "accuracy"]
MEAN_LOSS_METRICS = {"bias", "mae", "rmse", "brier", "logloss"}
# taken as the reference when truth is not given (task 'auto' or 'binary'); with none, task 'auto' reads the model
# columns as prediction errors
TRUTH_VARIANTS = ['truth', 'reference', 'reference_standard', 'ground_truth', 'groundtruth', 'label', 'y', 'y_true',
                  'target', 'diagnosis', 'outcome', 'actual', 'observed', 'measured', 'gold_standard']
# numeric columns never taken as models (subgroup columns are excluded too)
GENERIC_IGNORE = ['axl', 'al', 'axial_length', 'axiallength', 'acd', 'acdepth', 'k1', 'k2', 'meank', 'mean_k', 'k',
                  'km', 'kmean', 'k_mean', 'avg_k', 'lt', 'lensthickness', 'cct',
                  'wtw', 'hcd', 'iol_implanted', 'postoperative_refraction', 'corneal_refractive_index', 'eye', 'side',
                  'laterality', 'od_os', 're_le', 'age', 'gender', 'gender_num', 'sex', 'visit',
                  'site', 'centre', 'center', 'device', 'image_id', 'imageid']
SUBGROUP_MAX_LEVELS = 10   # a numeric subgroup with more distinct values is split into tertiles by default
ADJUSTED_CI_METHODS = ('romano_wolf', 'bonferroni', 'holm', 'hommel', 'fdr_bh')   # methods with adjusted intervals
P_ADJ_LABELS = {'romano_wolf': 'Romano-Wolf step-down', 'bonferroni': 'Bonferroni', 'holm': 'Holm',
                'hommel': 'Hommel', 'fdr_bh': 'Benjamini-Hochberg (FDR)', 'fdr_by': 'Benjamini-Yekutieli (FDR)',
                'sidak': 'Sidak', 'holm-sidak': 'Holm-Sidak', 'simes-hochberg': 'Hochberg',
                'fdr_tsbh': 'two-stage Benjamini-Hochberg (FDR)',
                'fdr_tsbky': 'two-stage Benjamini-Krieger-Yekutieli (FDR)'}
SINGLE_P = "single comparison (p_adj = p)"
SINGLE_CI = "single comparison (CI adj = CI)"

DEFAULT_SETTINGS = {
    "bootstrap_iterations": BOOTSTRAP_ITERATIONS,  # resamples (p-values on a grid of 1/(B+1))
    "seed": None,                  # None: BOOTSTRAP_SEED
    "fwer_method": MULTIPLE_TEST_METHOD,   # or 'bonferroni', 'holm', 'hommel', 'fdr_bh', 'none'
    "report_ci": True,             # confidence intervals for every estimate and difference
    "ci_level": CI_LEVEL,
    "within": MCNEMAR_THRESHOLD,   # continuous: tolerance for the proportion of eyes with |error| <= within
    "threshold": 0.5,              # binary: probability threshold for sensitivity, specificity and accuracy
    "invalid_threshold": None,     # continuous: exclude eyes with |error| above this
    "weighting": "eye",            # mean-loss metrics: 'eye', 'patient' or 'optimal' (minimum variance)
    "unit": "patient",             # 'patient' (recommended) or 'eye' (every eye independent; for comparison)
    "one_eye": False,              # keep one randomly chosen eye per patient
    "clip": 1e-15,                 # probabilities clipped to [clip, 1 - clip] for log-loss
    "metrics": None,               # subset of the metrics for the task (default: all)
    "min_group_size": MIN_GROUP_SIZE,
    "subgroup_cuts": None,         # numeric subgroup column: cut points, e.g. [22, 24.5] -> <22, 22 to <24.5, >=24.5
    "subgroup_bins": None,         # numeric subgroup column: number of equal-count quantile groups, e.g. 3 (tertiles)
    "cross_min_patients": 20,      # crossed subgroups: cells with fewer patients get point estimates only
    "verbose": True,
}


def _settings(settings):
    s = dict(DEFAULT_SETTINGS)
    s.update(settings or {})
    if s["seed"] is None:
        s["seed"] = BOOTSTRAP_SEED
    if s["weighting"] not in ("eye", "patient", "optimal"):
        raise ValueError("weighting must be 'eye', 'patient' or 'optimal'")
    if s["subgroup_cuts"] is not None and s["subgroup_bins"] is not None:
        raise ValueError("give subgroup_cuts or subgroup_bins, not both")
    if s["subgroup_cuts"] is not None:
        cuts = sorted({float(c) for c in np.atleast_1d(s["subgroup_cuts"])})
        if not cuts or not np.all(np.isfinite(cuts)):
            raise ValueError("subgroup_cuts must be one or more finite numbers")
        s["subgroup_cuts"] = cuts
    if s["subgroup_bins"] is not None:
        if int(s["subgroup_bins"]) != s["subgroup_bins"] or int(s["subgroup_bins"]) < 2:
            raise ValueError("subgroup_bins must be an integer of at least 2")
        s["subgroup_bins"] = int(s["subgroup_bins"])
    return s


# ----------------------------------------------------------------------------------------------------------------------
# Tests
# ----------------------------------------------------------------------------------------------------------------------

def _calculate_percentile_bootstrap_p(boot_diffs):
    """Two-sided percentile bootstrap p-value; ties (difference 0) count in both tails."""
    valid_diffs = boot_diffs[~np.isnan(boot_diffs)]
    if len(valid_diffs) == 0:
        return 1.0
    tail = min(np.sum(valid_diffs >= 0), np.sum(valid_diffs <= 0))
    p = 2 * (1 + tail) / (len(valid_diffs) + 1)
    return min(p, 1.0)


def _bootstrap_p(null_stats, obs_stat):
    """Bootstrap p-value (1 + #{null >= obs}) / (B + 1); never exactly 0 (Davison & Hinkley 1997)."""
    return (1 + np.sum(null_stats >= obs_stat)) / (len(null_stats) + 1)


def mcnemar_test(mask1, mask2):
    """McNemar test with continuity correction (one eye per patient)."""
    b = np.sum(mask1 & ~mask2)
    c = np.sum(~mask1 & mask2)
    n = b + c
    if b == c:
        return 1.0
    chi2_stat = ((abs(b - c) - 1.0) ** 2) / n
    p_val = stats.chi2.sf(chi2_stat, 1)
    return min(1.0, float(p_val))


def clustered_mcnemar_test(mask1, mask2, cluster_ids):
    """Yang, Sun & Hardin (2010) modified Obuchowski test for bilateral (clustered) data.

    D_k = b_k - c_k per patient (n_k eyes):
        X2 = ((K-1)/K) * (sum D_k)^2 / (0.5 * sum[(D_k - n_k*(p1 - p2))^2 + D_k^2])
    referred to chi-square with 1 df (no continuity correction).
    """
    d = mask1.astype(int) - mask2.astype(int)
    unique_clusters, cluster_idx = np.unique(cluster_ids, return_inverse=True)
    K = len(unique_clusters)
    if K < 2:
        return 1.0

    D_k = np.bincount(cluster_idx, weights=d)
    n_k = np.bincount(cluster_idx)
    sum_D = np.sum(D_k)
    if sum_D == 0:
        return 1.0

    p_diff = sum_D / np.sum(n_k)
    denom = 0.5 * np.sum((D_k - n_k * p_diff)**2 + D_k**2)
    if denom == 0:
        return 1.0

    chi2_stat = ((K - 1) / K) * sum_D**2 / denom
    return min(1.0, float(stats.chi2.sf(chi2_stat, 1)))


def wild_cluster_sd_all_pairs(X, cluster_idx, K, W_g, pairs=None, chunk=500):
    """Morgan-Pitman test of equal SDs with a wild cluster bootstrap-t, for every pair of models at once.

    The standardized difference and sum of the two errors are regressed on each other and the slope tested with a
    null-restricted wild cluster bootstrap (Cameron, Gelbach & Miller 2008; Rademacher weights W_g, B x K; CR0
    standard error). With weights of ±1 every bootstrap statistic is a function of cluster sums, so all resamples
    and pairs come from one matrix product. Returns {(j, k): (p, |t_obs|, |t*|)}."""
    X = np.asarray(X, float)
    n, J = X.shape
    B = W_g.shape[0]
    pairs = pairs or list(combinations(range(J), 2))
    out, cols, meta = {}, [], {}
    for (j, k) in pairs:
        diff, summ = X[:, j] - X[:, k], X[:, j] + X[:, k]
        var_d, var_s = np.var(diff, ddof=1), np.var(summ, ddof=1)
        if var_d == 0 or var_s == 0 or K < 3:
            out[(j, k)] = (1.0, 0.0, np.zeros(B)); continue
        x = (diff - np.mean(diff)) / np.sqrt(var_d)
        y = (summ - np.mean(summ)) / np.sqrt(var_s)
        x_bar, y_bar = np.mean(x), np.mean(y)
        dx = x - x_bar
        sxx = np.sum(dx ** 2)
        if sxx == 0:
            out[(j, k)] = (1.0, 0.0, np.zeros(B)); continue
        beta_hat = np.sum(dx * y) / sxx
        eps_hat = y - (y_bar + beta_hat * dx)
        var_b = np.sum(np.bincount(cluster_idx, weights=dx * eps_hat, minlength=K) ** 2) / sxx ** 2
        if var_b == 0:
            out[(j, k)] = (1.0, 0.0, np.zeros(B)); continue
        t_obs = beta_hat / np.sqrt(var_b)
        e_t = y - y_bar                                                     # null-restricted residuals
        A = np.bincount(cluster_idx, weights=dx * e_t, minlength=K)        # per-cluster sums
        E = np.bincount(cluster_idx, weights=e_t, minlength=K)
        D = np.bincount(cluster_idx, weights=dx, minlength=K)
        Q = np.bincount(cluster_idx, weights=dx ** 2, minlength=K)
        meta[(j, k)] = dict(c=len(cols), t_obs=t_obs, y_bar=y_bar, sxx=sxx, sdx=np.sum(dx),
                            DD=np.sum(D * D), AA=np.sum(A * A), QQ=np.sum(Q * Q), DQ=np.sum(D * Q))
        cols += [A, E, D * A, A * Q]
    if cols:
        V = np.column_stack(cols)
        G = np.empty((B, V.shape[1]))
        for i in range(0, B, chunk):
            G[i:i + chunk] = W_g[i:i + chunk].astype(float) @ V
        for key, m in meta.items():
            c = m["c"]
            wA, wE, wDA, wAQ = G[:, c], G[:, c + 1], G[:, c + 2], G[:, c + 3]
            beta_s = (m["y_bar"] * m["sdx"] + wA) / m["sxx"]                # bootstrap slope
            a = -wE / n                                                      # y_bar - bootstrap mean of y
            ss = (a ** 2 * m["DD"] + m["AA"] + beta_s ** 2 * m["QQ"]
                  + 2 * a * wDA - 2 * a * beta_s * m["DQ"] - 2 * beta_s * wAQ)  # sum of squared cluster scores
            var_s = np.maximum(ss, 0) / m["sxx"] ** 2
            t_s = np.zeros(B); ok = var_s > 0
            t_s[ok] = beta_s[ok] / np.sqrt(var_s[ok])
            out[key] = (float(_bootstrap_p(np.abs(t_s), np.abs(m["t_obs"]))), np.abs(m["t_obs"]), np.abs(t_s))
    return out


def comdvar(x, y):
    """Morgan-Pitman test of equal SDs with HC4 standard errors (one eye per patient).
    Adapted from Wilcox's comdvar (Rallfun v45)."""
    diff = x - y
    summ = x + y
    var_x, var_y = np.var(diff, ddof=1), np.var(summ, ddof=1)
    if var_x == 0 or var_y == 0:
        return 1.0
    z1 = (diff - np.mean(diff)) / np.sqrt(var_x)
    z2 = (summ - np.mean(summ)) / np.sqrt(var_y)
    return olshc4(z1, z2)


def olshc4(x, y):
    """OLS slope test with HC4 standard errors. Adapted from Wilcox's olshc4 (Rallfun v45)."""
    n = len(y)
    if n < 3:
        return 1.0

    X = np.column_stack((np.ones(n), x))
    try:
        XtX_inv = np.linalg.inv(X.T @ X)
    except np.linalg.LinAlgError:
        return 1.0

    h = np.sum((X @ XtX_inv) * X, axis=1)
    beta = XtX_inv @ X.T @ y
    res = y - X @ beta

    sum_h = np.sum(h)
    d = (n * h) / sum_h if sum_h != 0 else np.zeros_like(h)
    d = np.minimum(4, d)

    denom = np.power(1 - h, d)
    denom[denom < 1e-10] = 1e-10
    omega = (res**2) / denom

    middle = (X.T * omega) @ X
    Sigma = XtX_inv @ middle @ XtX_inv

    se_slope = np.sqrt(Sigma[1, 1])
    if se_slope == 0:
        return 1.0

    t_stat = beta[1] / se_slope
    df = n - 2
    return 2 * (1 - stats.t.cdf(np.abs(t_stat), df))


# ----------------------------------------------------------------------------------------------------------------------
# Patient bootstrap engines: every pair of models at once
# ----------------------------------------------------------------------------------------------------------------------

def _resample_counts(idx_chunk, K):
    """Count matrix: how often each cluster (patient or eye) appears in each resample of the block."""
    c = idx_chunk.shape[0]
    flat = (idx_chunk + (np.arange(c)[:, None] * K)).ravel()
    return np.bincount(flat, minlength=c * K).reshape(c, K).astype(np.float64)


def all_pairs_bootstrap(X, cluster_idx, K, b_indices, pairs=None, chunk=None, ci_level=None, prop_threshold=None):
    """Patient bootstrap tests of the mean error (MPE), MAE, RMSE and median absolute error (MedAE) for every pair of
    models. X: n x J prediction errors. Bootstrap-t with a CR1 standard error (RMSE tested on the MSE); MedAE by
    percentile bootstrap. Each block of resamples is a count matrix W (resamples x patients), so every bootstrap sum
    is one product W @ V, and each model's bootstrap median comes from running totals of W over its sorted errors.

    With ci_level, also the CIs of each difference from the same resamples: bootstrap-t for MPE and MAE, percentile
    for RMSE, MedAE, SD and (with prop_threshold) the proportion of eyes with |error| <= prop_threshold. RW_* hold
    the Romano-Wolf statistics. Returns {(j, k): result}.
    """
    X = np.asarray(X, float)
    n, J = X.shape
    B = b_indices.shape[0]
    pairs = pairs or list(combinations(range(J), 2))
    N_g = np.bincount(cluster_idx, minlength=K).astype(float)
    N_tot = N_g.sum()

    # per-patient sums for each model: absolute, squared and signed error
    fam = {"MAE": np.abs(X), "MSE": X ** 2, "MPE": X}
    S = {f: np.stack([np.bincount(cluster_idx, weights=v[:, j], minlength=K) for j in range(J)], axis=1) for f, v in fam.items()}
    means = {f: S[f].sum(0) / N_tot for f in fam}
    C = {f: S[f] - np.outer(N_g, means[f]) for f in fam}               # centered sums (precision)

    # columns of V: centered sums and their cross-products, for each family
    tri = [(a, b) for a in range(J) for b in range(a, J)]
    cols, index = [N_g], {}
    for f in fam:
        index[(f, "s")] = len(cols); cols += [C[f][:, j] for j in range(J)]
        index[(f, "x")] = len(cols); cols += [C[f][:, a] * C[f][:, b] for a, b in tri]
        index[(f, "n")] = len(cols); cols += [C[f][:, j] * N_g for j in range(J)]
    index["nn"] = len(cols); cols.append(N_g ** 2)
    index[("MSE", "raw")] = len(cols); cols += [S["MSE"][:, j] for j in range(J)]   # for RMSE differences
    if ci_level:
        index["SQ"] = len(cols)                                                    # sums of squared deviations (SD)
        cols += [np.bincount(cluster_idx, weights=(X[:, j] - means["MPE"][j]) ** 2, minlength=K) for j in range(J)]
        if prop_threshold is not None:
            index["PROP"] = len(cols)                                              # eyes with |PE| <= threshold
            cols += [np.bincount(cluster_idx, weights=(np.abs(X[:, j]) <= prop_threshold).astype(float), minlength=K) for j in range(J)]
    V = np.column_stack(cols)
    tpos = {t: i for i, t in enumerate(tri)}

    def cr1_se(Svec, mean_val):
        res = Svec - mean_val * N_g
        return np.sqrt(max(0.0, (K / (K - 1)) * np.sum(res ** 2) / N_tot ** 2)) if K > 1 else 0.0

    # medians: one per model per resample, from running totals of the counts over each model's sorted values
    absX = np.abs(X)
    order = [np.argsort(absX[:, j], kind="stable") for j in range(J)]
    sorted_vals = [absX[order[j], j] for j in range(J)]
    eye_cluster = [cluster_idx[order[j]] for j in range(J)]

    # A resample's median almost always lies close to the middle of the original sorted errors, so the running
    # totals are only needed over a window of sorted positions around the middle (+/- 5 sqrt(n) + 16 eyes): the eyes
    # drawn below the window come from one product with per-cluster counts. A resample whose middle rank falls outside
    # the window is redone over all the values, so every median is exactly the one the full calculation gives.
    n_eyes = len(cluster_idx)
    half = int(np.ceil(MEDIAN_WINDOW * np.sqrt(n_eyes))) + 16
    w_start, w_end = max(0, n_eyes // 2 - half), min(n_eyes, n_eyes // 2 + half + 1)
    windowed = (w_end - w_start) < n_eyes
    if windowed:
        below = [np.bincount(eye_cluster[j][:w_start], minlength=K).astype(float) for j in range(J)]
        in_window = [eye_cluster[j][w_start:w_end] for j in range(J)]

    def _positions(Wj, lo, hi, base=None):
        """Sorted positions of the lo-th and hi-th drawn eye (running totals over the columns of Wj)."""
        c = Wj.shape[0]
        cum = np.cumsum(Wj, axis=1)                                      # eyes drawn up to each sorted value
        if base is not None:
            cum += base[:, None]
        big = int(max(cum[:, -1].max(), hi.max())) + 1
        offs = np.arange(c) * big
        flat = (cum + offs[:, None]).ravel()
        start = np.arange(c) * cum.shape[1]
        pos_lo = np.searchsorted(flat, lo + offs, side="left") - start
        pos_hi = np.searchsorted(flat, hi + offs, side="left") - start
        return pos_lo, pos_hi

    def _medians(W):
        c = W.shape[0]
        out = np.empty((c, J))
        nb = W @ N_g                                                     # eyes in each resample
        lo = np.floor((nb + 1) / 2).astype(np.int64)                     # ranks of the middle value(s), 1-based
        hi = np.floor(nb / 2).astype(np.int64) + 1
        for j in range(J):
            if windowed:
                base = W @ below[j]                                      # eyes drawn below the window
                Wj = W[:, in_window[j]]
                pos_lo, pos_hi = _positions(Wj, lo, hi, base)
                inside = (base < lo) & (base + Wj.sum(axis=1) >= hi)
                pos_lo = np.clip(pos_lo, 0, w_end - w_start - 1) + w_start
                pos_hi = np.clip(pos_hi, 0, w_end - w_start - 1) + w_start
                if not inside.all():                                     # rare: redo these resamples in full
                    redo = ~inside
                    pos_lo[redo], pos_hi[redo] = _positions(W[redo][:, eye_cluster[j]], lo[redo], hi[redo])
            else:
                pos_lo, pos_hi = _positions(W[:, eye_cluster[j]], lo, hi)
            out[:, j] = (sorted_vals[j][pos_lo] + sorted_vals[j][pos_hi]) / 2
        return out

    chunk = chunk or max(1, min(500, int(2e7 // max(1, int(N_tot)))))
    sums = np.empty((B, V.shape[1]))
    med = np.empty((B, J))
    for i in range(0, B, chunk):
        idx = b_indices[i:i + chunk]
        W = _resample_counts(idx, K)
        sums[i:i + chunk] = W @ V
        med[i:i + chunk] = _medians(W)

    bN = sums[:, 0]
    a_ci = (1 - ci_level) / 2 if ci_level else None
    qci = lambda v: tuple(float(x) for x in np.quantile(v, [a_ci, 1 - a_ci]))
    out = {}
    for (j, k) in pairs:
        r = {}
        for f, key in (("MPE", "MPE"), ("MAE", "MAE"), ("MSE", "RMSE")):
            d_mean_obs = means[f][j] - means[f][k]
            Sd = S[f][:, j] - S[f][:, k]
            se_obs = cr1_se(Sd, d_mean_obs)
            t_obs = d_mean_obs / se_obs if se_obs > 0 else 0.0
            # bootstrap: centred difference sums and CR1 variance from the matrix products
            s = index[(f, "s")]; x = index[(f, "x")]; nn_ = index[(f, "n")]
            sd = sums[:, s + j] - sums[:, s + k]                                    # sum_g w (Cj - Ck)
            sdd = sums[:, x + tpos[(j, j)]] + sums[:, x + tpos[(k, k)]] - 2 * sums[:, x + tpos[(min(j, k), max(j, k))]]
            sdn = sums[:, nn_ + j] - sums[:, nn_ + k]
            snn = sums[:, index["nn"]]
            m = sd / bN                                                              # boot mean minus observed mean
            ssq = sdd - 2 * m * sdn + m ** 2 * snn                                   # sum_g w (S'_g - m N_g)^2
            var = (K / (K - 1)) * np.maximum(ssq, 0) / bN ** 2 if K > 1 else np.zeros(B)
            se = np.sqrt(var)
            with np.errstate(divide="ignore", invalid="ignore"):
                t_star = np.where(se > 0, m / se, 0.0)
            if ci_level:
                if key == "RMSE":                                                    # percentile, on the RMSE scale
                    q = index[("MSE", "raw")]
                    r[f"CI_{key}"] = qci(np.sqrt(np.maximum(0, sums[:, q + j] / bN)) - np.sqrt(np.maximum(0, sums[:, q + k] / bN)))
                else:                                                                # bootstrap-t
                    lo_q, hi_q = qci(t_star)
                    r[f"CI_{key}"] = (d_mean_obs - hi_q * se_obs, d_mean_obs - lo_q * se_obs)
            r[f"P_{key}"] = _bootstrap_p(np.abs(t_star), np.abs(t_obs))
            r[f"RW_{key}_obs"] = np.abs(t_obs); r[f"RW_{key}_null"] = np.abs(t_star)
            # RMSE is tested on the MSE scale, so its statistic gives no interval for the RMSE difference
            r[f"RW_{key}_diff"] = np.nan if key == "RMSE" else d_mean_obs
            r[f"RW_{key}_scale"] = np.nan if key == "RMSE" else se_obs
        boot_medae = med[:, j] - med[:, k]
        obs_medae = np.median(absX[:, j]) - np.median(absX[:, k])
        r["P_MedAE"] = _calculate_percentile_bootstrap_p(boot_medae)
        if ci_level:
            r["CI_MedAE"] = qci(boot_medae)
            s0, sq = index[("MPE", "s")], index["SQ"]
            with np.errstate(divide="ignore", invalid="ignore"):
                sd_j = np.sqrt(np.maximum(0, (sums[:, sq + j] - sums[:, s0 + j] ** 2 / bN) / (bN - 1)))
                sd_k = np.sqrt(np.maximum(0, (sums[:, sq + k] - sums[:, s0 + k] ** 2 / bN) / (bN - 1)))
            r["CI_SD"] = qci(sd_j - sd_k)
            if prop_threshold is not None:
                pc = index["PROP"]
                r["CI_PROP"] = qci((sums[:, pc + j] - sums[:, pc + k]) / bN)
        se_m = max(np.std(boot_medae, ddof=1), 1e-6)
        r["RW_MedAE_obs"] = abs(obs_medae) / se_m; r["RW_MedAE_null"] = np.abs(boot_medae - obs_medae) / se_m
        r["RW_MedAE_scale"] = se_m
        r["RW_MedAE_diff"] = obs_medae
        out[(j, k)] = r
    return out


def all_pairs_loss_bootstrap(L, cluster_idx, K, b_indices, weights=None, pairs=None, ci_level=None,
                             sqrt_keys=(), extra=None, chunk=None):
    """Patient bootstrap-t (CR1) for differences in the (weighted) mean of any per-eye loss, every pair of models.

    L maps a name to an n x J matrix of losses; with unit weights the results equal all_pairs_bootstrap's. Keys in
    sqrt_keys are tested on the mean loss with a percentile CI on the root scale (RMSE), whose draws are returned
    as DRAWS_<key>. extra maps names to per-patient vectors whose bootstrap sums are returned (proportion CIs).
    Returns ({(j, k): result}, {name: bootstrap sums}).
    """
    names = list(L)
    n, J = L[names[0]].shape
    B = b_indices.shape[0]
    pairs = pairs or list(combinations(range(J), 2))
    w = np.ones(n) if weights is None else np.asarray(weights, float)
    N_g = np.bincount(cluster_idx, weights=w, minlength=K).astype(float)
    N_tot = N_g.sum()
    S = {f: np.stack([np.bincount(cluster_idx, weights=w * L[f][:, j], minlength=K) for j in range(J)], axis=1)
         for f in names}
    means = {f: S[f].sum(0) / N_tot for f in names}
    C = {f: S[f] - np.outer(N_g, means[f]) for f in names}
    tri = [(a, b) for a in range(J) for b in range(a, J)]
    cols, index = [N_g], {}
    for f in names:
        index[(f, "s")] = len(cols); cols += [C[f][:, j] for j in range(J)]
        index[(f, "x")] = len(cols); cols += [C[f][:, a] * C[f][:, b] for a, b in tri]
        index[(f, "n")] = len(cols); cols += [C[f][:, j] * N_g for j in range(J)]
    index["nn"] = len(cols); cols.append(N_g ** 2)
    for f in sqrt_keys:
        index[(f, "raw")] = len(cols); cols += [S[f][:, j] for j in range(J)]
    extra = extra or {}
    for name, v in extra.items():
        index[("extra", name)] = len(cols); cols.append(np.asarray(v, float))
    V = np.column_stack(cols)
    tpos = {t: i for i, t in enumerate(tri)}

    chunk = chunk or max(1, min(500, int(2e7 // max(1, n))))
    sums = np.empty((B, V.shape[1]))
    for i in range(0, B, chunk):
        sums[i:i + chunk] = _resample_counts(b_indices[i:i + chunk], K) @ V
    bN = sums[:, 0]
    a_ci = (1 - ci_level) / 2 if ci_level else None
    qci = lambda v: tuple(float(x) for x in np.quantile(v[np.isfinite(v)], [a_ci, 1 - a_ci]))

    def cr1_se(Svec, mean_val):
        res = Svec - mean_val * N_g
        return np.sqrt(max(0.0, (K / (K - 1)) * np.sum(res ** 2) / N_tot ** 2)) if K > 1 else 0.0

    out = {}
    for (j, k) in pairs:
        r = {}
        for f in names:
            d_mean_obs = means[f][j] - means[f][k]
            se_obs = cr1_se(S[f][:, j] - S[f][:, k], d_mean_obs)
            t_obs = d_mean_obs / se_obs if se_obs > 0 else 0.0
            s = index[(f, "s")]; x = index[(f, "x")]; nn_ = index[(f, "n")]
            sd = sums[:, s + j] - sums[:, s + k]
            sdd = sums[:, x + tpos[(j, j)]] + sums[:, x + tpos[(k, k)]] - 2 * sums[:, x + tpos[(min(j, k), max(j, k))]]
            sdn = sums[:, nn_ + j] - sums[:, nn_ + k]
            snn = sums[:, index["nn"]]
            m = sd / bN
            ssq = sdd - 2 * m * sdn + m ** 2 * snn
            var = (K / (K - 1)) * np.maximum(ssq, 0) / bN ** 2 if K > 1 else np.zeros(B)
            se = np.sqrt(var)
            with np.errstate(divide="ignore", invalid="ignore"):
                t_star = np.where(se > 0, m / se, 0.0)
            r[f"P_{f}"] = _bootstrap_p(np.abs(t_star), np.abs(t_obs)) if se_obs > 0 else 1.0
            r[f"RW_{f}_obs"] = np.abs(t_obs); r[f"RW_{f}_null"] = np.abs(t_star)
            # root-scale keys (RMSE from squared errors) are tested on the mean loss: no adjusted interval on the root scale
            r[f"RW_{f}_diff"] = np.nan if f in sqrt_keys else d_mean_obs
            r[f"RW_{f}_scale"] = np.nan if f in sqrt_keys else se_obs
            r[f"DIFF_{f}"] = d_mean_obs
            r[f"EST_{f}"] = (means[f][j], means[f][k])
            if ci_level:
                if f in sqrt_keys:
                    q = index[(f, "raw")]
                    draws = np.sqrt(np.maximum(0, sums[:, q + j] / bN)) - np.sqrt(np.maximum(0, sums[:, q + k] / bN))
                    r[f"CI_{f}"] = qci(draws)
                    r[f"DRAWS_{f}"] = draws
                else:
                    lo_q, hi_q = qci(t_star)
                    r[f"CI_{f}"] = (d_mean_obs - hi_q * se_obs, d_mean_obs - lo_q * se_obs)
        out[(j, k)] = r
    extra_sums = {name: sums[:, index[("extra", name)]] for name in extra}
    return out, extra_sums


def _resampled_cluster_sums(cols, b_indices, K, chunk=500):
    """B x C matrix: the sum of each per-cluster vector in cols over the patients drawn in each resample."""
    V = np.column_stack(cols)
    out = np.empty((b_indices.shape[0], V.shape[1]))
    for i in range(0, b_indices.shape[0], chunk):
        out[i:i + chunk] = _resample_counts(b_indices[i:i + chunk], K) @ V
    return out


def continuous_percentile_draws(E, cluster_idx, K, b_indices, within, pairs=None):
    """Bootstrap draws of the RMSE, SD and proportion-within differences for every pair of models (as behind the
    percentile CIs of all_pairs_bootstrap), for the Bonferroni-type adjusted intervals."""
    E = np.asarray(E, float)
    n, J = E.shape
    pairs = pairs or list(combinations(range(J), 2))
    N_g = np.bincount(cluster_idx, minlength=K).astype(float)
    S = np.stack([np.bincount(cluster_idx, weights=E[:, j], minlength=K) for j in range(J)], axis=1)
    mean = S.sum(0) / N_g.sum()
    C = S - np.outer(N_g, mean)
    sq = [np.bincount(cluster_idx, weights=(E[:, j] - mean[j]) ** 2, minlength=K) for j in range(J)]
    raw = [np.bincount(cluster_idx, weights=(E ** 2)[:, j], minlength=K) for j in range(J)]
    prop = [np.bincount(cluster_idx, weights=(np.abs(E[:, j]) <= within).astype(float), minlength=K) for j in range(J)]
    sums = _resampled_cluster_sums([N_g] + [C[:, j] for j in range(J)] + sq + raw + prop, b_indices, K)
    bN = sums[:, :1]
    c0, s0, r0, p0 = 1, 1 + J, 1 + 2 * J, 1 + 3 * J
    with np.errstate(divide="ignore", invalid="ignore"):
        sd = np.sqrt(np.maximum(0, (sums[:, s0:s0 + J] - sums[:, c0:c0 + J] ** 2 / bN) / (bN - 1)))
    rmse = np.sqrt(np.maximum(0, sums[:, r0:r0 + J] / bN))
    pw = sums[:, p0:p0 + J] / bN
    return {"rmse": {(j, k): rmse[:, j] - rmse[:, k] for j, k in pairs},
            "sd": {(j, k): sd[:, j] - sd[:, k] for j, k in pairs},
            "within": {(j, k): pw[:, j] - pw[:, k] for j, k in pairs}}


# ----------------------------------------------------------------------------------------------------------------------
# AUC: Obuchowski (1997) clustered ROC; DeLong et al (1988) when every patient has one eye
# ----------------------------------------------------------------------------------------------------------------------

def _placements(x_pos, y_neg):
    ys, xs = np.sort(y_neg), np.sort(x_pos)
    lt = np.searchsorted(ys, x_pos, side="left"); le = np.searchsorted(ys, x_pos, side="right")
    v10 = (lt + 0.5 * (le - lt)) / len(y_neg)
    gt = len(x_pos) - np.searchsorted(xs, y_neg, side="right"); ge = len(x_pos) - np.searchsorted(xs, y_neg, side="left")
    v01 = (gt + 0.5 * (ge - gt)) / len(x_pos)
    return v10, v01


def auc_all_pairs(P, y, cluster_idx, K, pairs=None, ci_level=0.95):
    """Nonparametric clustered AUC and its comparison between models (Obuchowski 1997, equations 3-5).

    P: n x J predicted probabilities (or any scores); y: 0/1 reference standard per eye. The covariance between
    two models' AUCs uses the patient as the unit, so both eyes can be kept; with one eye per patient the method
    reduces to DeLong et al. (1988). Returns {(j, k): (auc_j, auc_k, diff, ci_low, ci_high, p, se)}."""
    y = np.asarray(y).astype(int)
    n, J = P.shape
    pos, neg = y == 1, y == 0
    mi = np.bincount(cluster_idx, weights=pos, minlength=K)
    ni = np.bincount(cluster_idx, weights=neg, minlength=K)
    M, Nn = mi.sum(), ni.sum()
    I10, I01 = int(np.sum(mi > 0)), int(np.sum(ni > 0))
    pairs = pairs or list(combinations(range(J), 2))
    out = {}
    if M == 0 or Nn == 0 or I10 < 2 or I01 < 2:
        return {pr: (np.nan, np.nan, np.nan, np.nan, np.nan, 1.0, np.nan) for pr in pairs}
    th, r10, r01 = [], [], []
    for j in range(J):
        v10, v01 = _placements(P[pos, j], P[neg, j])
        t = float(np.mean(v10))
        th.append(t)
        r10.append(np.bincount(cluster_idx[pos], weights=v10, minlength=K) - mi * t)
        r01.append(np.bincount(cluster_idx[neg], weights=v01, minlength=K) - ni * t)
    S10 = lambda a, b: I10 / ((I10 - 1) * M) * np.sum(r10[a] * r10[b])
    S01 = lambda a, b: I01 / ((I01 - 1) * Nn) * np.sum(r01[a] * r01[b])
    S11 = lambda a, b: K / (K - 1) * np.sum(r10[a] * r01[b])
    z = stats.norm.ppf(1 - (1 - (ci_level or 0.95)) / 2)
    for (j, k) in pairs:
        var = [S10(r, r) / M + S01(r, r) / Nn + 2 * S11(r, r) / (M * Nn) for r in (j, k)]
        cov = S10(j, k) / M + S01(j, k) / Nn + (S11(j, k) + S11(k, j)) / (M * Nn)
        se = float(np.sqrt(max(var[0] + var[1] - 2 * cov, 0.0)))
        diff = th[j] - th[k]
        p = float(2 * stats.norm.sf(abs(diff) / se)) if se > 0 else 1.0
        out[(j, k)] = (th[j], th[k], diff, diff - z * se, diff + z * se, p, se)
    return out


# ----------------------------------------------------------------------------------------------------------------------
# Fellow-eye correlation and the value of a second eye
# ----------------------------------------------------------------------------------------------------------------------

def fellow_eye_correlation(values, cluster_idx):
    """Pairwise intraclass correlation of a per-eye quantity between eyes of the same patient (rho).

    Applied to the per-eye difference in loss between two models it gives ρΔ, the quantity that decides
    how much eye-level tests are misled and how much a second eye adds (Kerry & Bland 2001; Eldridge 2006)."""
    values = np.asarray(values, float)
    m = np.bincount(cluster_idx).astype(float)
    a = np.bincount(cluster_idx, weights=values)
    q = np.bincount(cluster_idx, weights=values ** 2)
    N, NP = m.sum(), np.sum(m * (m - 1))
    if NP == 0 or N == 0:
        return np.nan
    mu = a.sum() / N
    s2 = q.sum() / N - mu ** 2
    if s2 <= 0:
        return np.nan
    return float((np.sum(a ** 2 - q) / NP - mu ** 2) / s2)


def one_eye_factor(sizes, rho, weighting="eye"):
    """How many times more patients one eye per patient needs for the same precision as the analysis of all
    eyes with the given weighting (design effect for unequal cluster sizes; Kerry & Bland 2001)."""
    m = np.asarray(sizes, float)
    if not np.isfinite(rho) or len(m) == 0:
        return np.nan
    r = float(np.clip(rho, 0.0, 0.999))
    K = len(m)
    var = {"eye": np.sum(m * (1 + (m - 1) * r)) / np.sum(m) ** 2,
           "patient": np.sum((1 + (m - 1) * r) / m) / K ** 2,
           "optimal": 1.0 / np.sum(m / (1 + (m - 1) * r))}[weighting]
    return float((1.0 / K) / var)


def eye_weights(cluster_idx, weighting, rho=0.0):
    """Per-eye weights: 'eye' 1; 'patient' 1/m (each patient counts once); 'optimal' 1/(1 + (m - 1) rho)."""
    m = np.bincount(cluster_idx).astype(float)[cluster_idx]
    if weighting == "eye":
        return np.ones(len(cluster_idx))
    if weighting == "patient":
        return 1.0 / m
    r = float(np.clip(rho if np.isfinite(rho) else 0.0, 0.0, 0.999))
    return 1.0 / (1.0 + (m - 1.0) * r)


# ----------------------------------------------------------------------------------------------------------------------
# Multiplicity: adjusted p-values and matching adjusted intervals
# ----------------------------------------------------------------------------------------------------------------------

def romano_wolf_correction(obs_stats, boot_null_stats):
    """Romano-Wolf step-down adjusted p-values from the bootstrap null distributions of |t|."""
    M = len(obs_stats)
    if M == 0:
        return np.array([])

    adj_p = np.zeros(M)
    order = np.argsort(obs_stats)[::-1]
    sorted_nulls = boot_null_stats[:, order]

    max_nulls = np.maximum.accumulate(sorted_nulls[:, ::-1], axis=1)[:, ::-1]

    for i in range(M):
        p_val = _bootstrap_p(max_nulls[:, i], obs_stats[order[i]])
        adj_p[order[i]] = p_val

    for i in range(1, M):
        adj_p[order[i]] = max(adj_p[order[i]], adj_p[order[i-1]])

    return adj_p


def adjusted_ci_critical(boot_null_stats, level, method, n_reject=None):
    """Critical value c of each comparison's adjusted interval (difference ± c × scale), matching the
    p-value adjustment. boot_null_stats: B x M bootstrap |t|.

    romano_wolf     one c: the level-quantile of the largest |t| across comparisons (single-step max-|t|)
    bonferroni      each comparison's 1 - (1 - level)/M quantile of |t|; also used for holm and hommel
    fdr_bh          Benjamini-Yekutieli: level 1 - R(1 - level)/M, R = comparisons with adjusted p <= 1 - level
    Returns M critical values, or None for other methods."""
    N = np.asarray(boot_null_stats, float)
    M, alpha = N.shape[1], 1 - level
    if method == 'romano_wolf':
        return np.full(M, np.quantile(N.max(axis=1), level))
    if method in ('bonferroni', 'holm', 'hommel'):
        q = 1 - alpha / M
    elif method == 'fdr_bh':
        q = 1 - max(int(n_reject or 0), 1) * alpha / M
    else:
        return None
    return np.quantile(N, q, axis=0)


def adjusted_ci_description(method, level):
    """One-line description of the adjusted intervals for the methodology sheet."""
    L = f"{level * 100:g}%"
    if method == 'romano_wolf':
        return (f"{L} simultaneous intervals: difference \u00b1 c \u00d7 SE, c = {L} quantile of the largest bootstrap |t| across "
                "comparisons (single-step max-|t|, the interval counterpart of Romano-Wolf)")
    if method == 'bonferroni':
        return (f"{L} Bonferroni intervals: difference \u00b1 c \u00d7 SE, c = 1 - {1 - level:g}/m quantile of each comparison's "
                "bootstrap |t| (m comparisons); they agree with the Bonferroni-adjusted p-values")
    if method in ('holm', 'hommel'):
        return (f"{L} Bonferroni intervals (c = 1 - {1 - level:g}/m quantile of each comparison's bootstrap |t|): "
                f"{method.capitalize()} has no compatible finite intervals, so these are conservative and an interval can "
                "include 0 next to a significant adjusted p-value")
    if method == 'fdr_bh':
        return (f"Benjamini-Yekutieli false-coverage-rate intervals at level 1 - R \u00d7 {1 - level:g}/m (R = comparisons "
                f"with adjusted p \u2264 {1 - level:g}); they control the false coverage rate among the significant "
                "comparisons, not simultaneous coverage")
    return "None: adjusted intervals are given for Romano-Wolf, Bonferroni, Holm, Hommel and Benjamini-Hochberg only"


RW_MAP = {"Boot (MPE) p": "MPE", "Boot (MAE) p": "MAE", "Boot (MedAE) p": "MedAE", "Boot (RMSE) p": "RMSE",
          "Morgan-Pitman (SD) p": "SD", "Boot (Brier) p": "Brier", "Boot (LogLoss) p": "LogLoss"}


def _adjust(wide, rw_stats, method, methods_out=None):
    """Adjusted p-values for every p column: Romano-Wolf where bootstrap null distributions exist (else Hommel), or
    the chosen method. methods_out, if given, receives the adjustment applied to each column."""
    if wide.empty:
        return wide
    methods_out = {} if methods_out is None else methods_out
    for col in [c for c in wide.columns if c.endswith(" p")]:
        pvals = wide[col].values.astype(float)
        adj = np.full(len(pvals), np.nan)
        mask = ~np.isnan(pvals)
        use_fallback = np.sum(mask) > 1
        if not use_fallback:
            adj[mask] = pvals[mask]
            methods_out[col] = SINGLE_P
        if method == "romano_wolf" and col in RW_MAP and use_fallback and RW_MAP[col] in rw_stats:
            obs = np.array(rw_stats[RW_MAP[col]]["obs"])
            nulls = rw_stats[RW_MAP[col]]["null"]
            if len(obs) > 1 and len(nulls) > 0:
                null_matrix = np.column_stack(nulls)
                if np.sum(np.abs(null_matrix)) > 0:
                    adj_rw = romano_wolf_correction(obs, null_matrix)
                    adj[mask] = adj_rw[mask]
                    use_fallback = False
                    methods_out[col] = P_ADJ_LABELS["romano_wolf"]
        if use_fallback and method not in (None, "none"):
            current = FALLBACK_TEST_METHOD if method == "romano_wolf" else method
            _, a, _, _ = multipletests(pvals[mask], method=current)
            adj[mask] = a
            methods_out[col] = P_ADJ_LABELS.get(current, str(current))
        elif use_fallback:
            adj[mask] = pvals[mask]
            methods_out[col] = "none (p_adj = p; fwer_method='none')"
        adj[mask] = np.maximum(adj[mask], pvals[mask])
        wide.insert(wide.columns.get_loc(col) + 1, col.replace(" p", " p_adj"), adj)
    # adjusted intervals next to each CI, matching the method used for the p-values (see adjusted_ci_critical)
    return add_adjusted_cis(wide, rw_stats, metrics=("MPE", "MAE", "MedAE", "RMSE", "Brier", "LogLoss"), method=method)


def add_adjusted_cis(df_results, rw_stats=None, metrics=('MPE', 'MAE', 'MedAE', 'RMSE'), method='romano_wolf'):
    """Adjusted ('CI adj') intervals next to each bootstrap-t interval, matching the multiplicity method
    (adjusted_ci_critical). With a single comparison the adjusted interval is the unadjusted one."""
    if df_results.empty or method not in ADJUSTED_CI_METHODS:
        return df_results
    for metric in metrics:
        dcol = f'{metric} Diff'
        lows = [c for c in df_results.columns if c.startswith(f'{dcol} ') and c.endswith('% CI low') and 'adj' not in c]
        if not lows:
            continue
        lbl = lows[0][len(dcol) + 1:-len(' low')]                       # e.g. '95% CI'
        level = float(lbl.split('%')[0]) / 100
        low_col, high_col = f'{dcol} {lbl} low', f'{dcol} {lbl} high'
        if len(df_results) == 1:
            adj_lo, adj_hi = df_results[low_col].values.astype(float), df_results[high_col].values.astype(float)
        else:
            st = (rw_stats or {}).get(metric, {})
            diffs, scales = np.array(st.get('diff', []), float), np.array(st.get('scale', []), float)
            if len(diffs) != len(df_results) or not (np.isfinite(diffs).all() and np.isfinite(scales).all()):
                continue                                                # no statistic to adjust with (e.g. RMSE)
            pcol = f'Boot ({metric}) p_adj'
            n_rej = int(np.sum(df_results[pcol].values <= 1 - level + 1e-12)) if pcol in df_results.columns else None
            c = adjusted_ci_critical(np.column_stack(st['null']), level, method, n_rej)
            adj_lo, adj_hi = diffs - c * scales, diffs + c * scales
        idx = df_results.columns.get_loc(high_col) + 1
        for name, vals in ((f'{dcol} {lbl} adj low', adj_lo), (f'{dcol} {lbl} adj high', adj_hi)):
            if name in df_results.columns:
                df_results[name] = vals
            else:
                df_results.insert(idx, name, vals)
                idx += 1
    return df_results


BOOTSTRAP_T_ADJ = {"bias", "mae", "medae", "brier", "logloss"}   # adjusted intervals from add_adjusted_cis
PERCENTILE_ADJ = {"rmse", "sd", "within", "sensitivity", "specificity", "accuracy"}   # Bonferroni percentile


def _adjusted_coverage(method, level, m, n_reject=0):
    """Per-comparison coverage of a Bonferroni-type adjusted interval, as in adjusted_ci_critical: 1 - α/m for
    Romano-Wolf (these metrics have no max-|t| statistic; their p-values use Hommel or a different statistic),
    Bonferroni, Holm and Hommel; Benjamini-Yekutieli 1 - Rα/m for fdr_bh (R = comparisons with p_adj <= α)."""
    alpha = 1 - level
    if method in ("romano_wolf", "bonferroni", "holm", "hommel"):
        return 1 - alpha / m
    if method == "fdr_bh":
        return 1 - max(int(n_reject), 1) * alpha / m
    return None


def _add_bonferroni_type_cis(wide, task, models, pct_draws, auc_se, s):
    """Bonferroni-type adjusted intervals (level q = 1 - α/m; 1 - Rα/m for fdr_bh) for the metrics
    without a bootstrap-t: AUC as difference ± z × SE; RMSE, SD and the proportions as percentile intervals of
    the same resamples. With a single pair they equal the unadjusted intervals."""
    if wide.empty or not s["report_ci"] or s["fwer_method"] not in ADJUSTED_CI_METHODS:
        return wide
    level = float(s["ci_level"])
    lbl = f"{level * 100:g}% CI"
    keys = [(models.index(a), models.index(b)) for a, b in zip(wide["Model A"], wide["Model B"])]
    for metric in (CONTINUOUS_METRICS if task == "continuous" else BINARY_METRICS):
        if metric not in PERCENTILE_ADJ and metric != "auc":
            continue
        dname, pname = WIDE_NAMES[metric]
        low_col, high_col = f"{dname} {lbl} low", f"{dname} {lbl} high"
        if low_col not in wide.columns or f"{dname} {lbl} adj low" in wide.columns:
            continue                                     # (add_adjusted_cis copies RMSE's CI for a single pair)
        if len(wide) == 1:
            adj_lo, adj_hi = wide[low_col].values.astype(float), wide[high_col].values.astype(float)
        else:
            src = auc_se if metric == "auc" else pct_draws.get(metric, {})
            m = int(np.sum(np.isfinite(wide[pname].values.astype(float))))
            n_rej = int(np.sum(wide[pname.replace(" p", " p_adj")].values <= 1 - level + 1e-12))
            q = _adjusted_coverage(s["fwer_method"], level, max(m, 1), n_rej)
            a = (1 - q) / 2
            adj_lo, adj_hi = np.full(len(wide), np.nan), np.full(len(wide), np.nan)
            for i, key in enumerate(keys):
                if key not in src or src[key] is None or not np.isfinite(wide[low_col].iloc[i]):
                    continue
                if metric == "auc":
                    d, se = src[key]
                    z = stats.norm.ppf(1 - a)
                    adj_lo[i], adj_hi[i] = d - z * se, d + z * se
                else:
                    v = np.asarray(src[key], float)
                    v = v[np.isfinite(v)]
                    if len(v):
                        adj_lo[i], adj_hi[i] = np.quantile(v, [a, 1 - a])
        idx = wide.columns.get_loc(high_col) + 1
        for name, vals in ((f"{dname} {lbl} adj low", adj_lo), (f"{dname} {lbl} adj high", adj_hi)):
            wide.insert(idx, name, vals)
            idx += 1
    return wide


def _adj_ci_matches(metric, fm):
    """True if this metric's adjusted interval matches its adjusted p-value: Romano-Wolf simultaneous
    intervals for the bootstrap-t metrics and MedAE, Bonferroni with Bonferroni, Benjamini-Yekutieli with
    Benjamini-Hochberg. Otherwise the CI stays unadjusted."""
    if fm in ("bonferroni", "fdr_bh"):
        return True
    return fm == "romano_wolf" and metric in BOOTSTRAP_T_ADJ


def _drop_unmatched_adjusted_cis(wide, task, s):
    """Remove the 'CI adj' columns of metrics whose adjusted p-value has no matching interval (_adj_ci_matches).
    With a single model pair nothing is adjusted and the columns stay (they equal the unadjusted intervals)."""
    if wide.empty or len(wide) == 1:
        return wide
    lbl = f"{float(s['ci_level']) * 100:g}% CI"
    drop = []
    for metric in (CONTINUOUS_METRICS if task == "continuous" else BINARY_METRICS):
        if not _adj_ci_matches(metric, s["fwer_method"]):
            drop += [c for c in (f"{WIDE_NAMES[metric][0]} {lbl} adj low", f"{WIDE_NAMES[metric][0]} {lbl} adj high")
                     if c in wide.columns]
    return wide.drop(columns=drop)


def _ci_adj_label(metric, wide, s, has_clusters=True):
    """How the adjusted interval of a metric was made (column ci_adj_method), or why there is none."""
    fm = s["fwer_method"]
    if (s["report_ci"] and fm in ADJUSTED_CI_METHODS and len(wide) > 1 and not _adj_ci_matches(metric, fm)):
        why = ("its Romano-Wolf p-value comes from a different statistic" if fm == "romano_wolf"
               and metric in ("rmse", "sd") else
               f"{'Hommel' if fm == 'romano_wolf' else P_ADJ_LABELS.get(fm, fm)} p-values have no matching interval")
        return f"not adjusted: {why} (the CI is unadjusted; fwer_method='bonferroni' gives matching intervals)"
    if not s["report_ci"]:
        return "not available: no confidence intervals (report_ci=False)"
    if fm in (None, "none"):
        return "not available: no multiplicity adjustment (fwer_method='none')"
    if fm not in ADJUSTED_CI_METHODS:
        return f"not available: {P_ADJ_LABELS.get(fm, fm)} has no interval counterpart"
    if f"{WIDE_NAMES[metric][0]} {float(s['ci_level']) * 100:g}% CI adj low" not in wide.columns:
        return "not available: no bootstrap statistic for this metric in this group"
    if len(wide) == 1:
        return SINGLE_CI
    if fm == "fdr_bh":
        head = "Benjamini-Yekutieli false-coverage-rate"
    elif metric in BOOTSTRAP_T_ADJ and fm == "romano_wolf":
        return "Romano-Wolf simultaneous (max-|t|)"
    else:
        head = "Bonferroni"
    if metric in BOOTSTRAP_T_ADJ:
        return f"{head} (bootstrap |t|)" if metric != "medae" else f"{head} (bootstrap, studentized by its SD)"
    if metric == "auc":
        return f"{head} (normal theory, {'Obuchowski' if has_clusters else 'DeLong'} SE)"
    return f"{head} (percentile bootstrap)"


# ----------------------------------------------------------------------------------------------------------------------
# Comparing models: one group of eyes
# ----------------------------------------------------------------------------------------------------------------------

def _detect_columns(df, task, models, patient, truth, extra_ignore=()):
    """Resolve the task and the patient, reference and model columns. Without truth, a column named as in
    TRUTH_VARIANTS is the reference (not looked for when task is 'continuous'). Without models, every numeric column
    other than the patient, reference, GENERIC_IGNORE and extra_ignore columns. 'auto' becomes 'binary' when the
    reference is 0/1 and every model value lies in [0, 1]."""
    cols_lower = {str(c).lower(): c for c in df.columns}
    if patient is None:
        patient = next((cols_lower[v] for v in ['patient_id', 'id_patient', 'patientid', 'pt_id', 'id', 'hosp_id',
                                                 'hospid', 'ptid', 'patient', 'subject', 'subject_id']
                        if v in cols_lower), None)
    if truth is None and task != "continuous":
        truth = next((cols_lower[v] for v in TRUTH_VARIANTS if v in cols_lower), None)
    if models is None:
        numeric = df.select_dtypes(include=[np.number]).columns
        skip = {str(patient).lower(), str(truth).lower()} | set(GENERIC_IGNORE) | {str(c).lower() for c in extra_ignore}
        models = [c for c in numeric if str(c).lower() not in skip and not str(c).lower().startswith('predict_')
                  and not str(c).lower().endswith('_const')]
        named_like_truth = [c for c in models if str(c).lower() in TRUTH_VARIANTS]
        if task == "continuous" and truth is None and named_like_truth:
            warnings.warn(f"column(s) {named_like_truth} are analyzed as prediction errors of a model; if one is the "
                          f"reference value, pass truth='{named_like_truth[0]}' (--truth) or use task='auto'",
                          stacklevel=3)
    if task == "auto":
        task = "continuous"
        if truth is not None:
            yv = pd.to_numeric(df[truth], errors="coerce").dropna()
            pv = df[models].stack()
            if set(np.unique(yv)) <= {0, 1} and pv.min() >= 0 and pv.max() <= 1:
                task = "binary"
    return task, list(models), patient, truth


def model_performance(P, y, task, models, metrics, cluster_idx, K, b_idx, s, has_clusters, group_name="All"):
    """Each model's own performance, every metric, with a percentile bootstrap CI that resamples patients (the same
    resamples as the comparisons), so fellow eyes stay together. Eyes are weighted equally."""
    n = P.shape[0]
    ci_level = float(s["ci_level"])
    a = (1 - ci_level) / 2
    method = "patient bootstrap, percentile" if has_clusters else "bootstrap, percentile"
    draws = {}                                   # metric -> bootstrap chunks
    if task == "continuous":
        E = P if y is None else P - y[:, None]
        absE = np.abs(E)
        within = (absE <= s["within"]).astype(float)
        order = np.argsort(absE, axis=0)
    else:
        clip = float(s["clip"])
        Pc = np.clip(P, clip, 1 - clip)
        brier = (P - y[:, None]) ** 2
        logloss = -(y[:, None] * np.log(Pc) + (1 - y[:, None]) * np.log(1 - Pc))
        correct = ((P >= s["threshold"]) == (y[:, None] == 1)).astype(float)
        pos, neg = (y == 1).astype(float), (y == 0).astype(float)
        auc_prep = []
        for j in range(P.shape[1]):
            o = np.argsort(P[:, j], kind="mergesort")
            sv = P[o, j]
            starts = np.r_[0, np.flatnonzero(np.diff(sv)) + 1]
            auc_prep.append((o, starts))
    point = {}
    for metric in metrics:
        if task == "continuous":
            point[metric] = {"bias": E.mean(0), "mae": absE.mean(0), "rmse": np.sqrt((E ** 2).mean(0)),
                             "medae": np.median(absE, 0), "sd": E.std(0, ddof=1) if n > 1 else np.full(E.shape[1], np.nan),
                             "within": within.mean(0)}.get(metric)
        else:
            if metric == "auc":
                pt = []
                for j in range(P.shape[1]):
                    xp, xn = P[y == 1, j], P[y == 0, j]
                    pt.append(_placements(xp, xn)[0].mean() if len(xp) and len(xn) else np.nan)
                point[metric] = np.array(pt)
            else:
                point[metric] = {"brier": brier.mean(0), "logloss": logloss.mean(0),
                                 "sensitivity": correct[y == 1].mean(0) if pos.sum() else np.full(P.shape[1], np.nan),
                                 "specificity": correct[y == 0].mean(0) if neg.sum() else np.full(P.shape[1], np.nan),
                                 "accuracy": correct.mean(0)}.get(metric)
        draws[metric] = []
    B = b_idx.shape[0]
    for c0 in range(0, B, 250):
        cnt = np.zeros((min(250, B - c0), K))
        for r, row in enumerate(b_idx[c0:c0 + 250]):
            cnt[r] = np.bincount(row, minlength=K)
        W = cnt[:, cluster_idx]                     # resample x eye weights (copies of each patient)
        tot = W.sum(1, keepdims=True)
        for metric in metrics:
            if task == "continuous":
                if metric == "bias":
                    v = (W @ E) / tot
                elif metric == "mae":
                    v = (W @ absE) / tot
                elif metric == "rmse":
                    v = np.sqrt((W @ (E ** 2)) / tot)
                elif metric == "within":
                    v = (W @ within) / tot
                elif metric == "sd":
                    m1 = (W @ E) / tot
                    v = np.sqrt(np.maximum((W @ (E ** 2)) - tot * m1 ** 2, 0) / np.maximum(tot - 1, 1))
                elif metric == "medae":
                    v = np.empty((W.shape[0], E.shape[1]))
                    for j in range(E.shape[1]):
                        cw = np.cumsum(W[:, order[:, j]], axis=1)
                        k = np.argmax(cw >= cw[:, -1:] / 2, axis=1)
                        v[:, j] = absE[order[:, j], j][k]
                else:
                    continue
            else:
                if metric == "brier":
                    v = (W @ brier) / tot
                elif metric == "logloss":
                    v = (W @ logloss) / tot
                elif metric in ("sensitivity", "specificity"):
                    sub = pos if metric == "sensitivity" else neg
                    with np.errstate(divide="ignore", invalid="ignore"):
                        v = (W @ (correct * sub[:, None])) / (W @ sub)[:, None]
                elif metric == "accuracy":
                    v = (W @ correct) / tot
                elif metric == "auc":
                    v = np.empty((W.shape[0], P.shape[1]))
                    for j, (o, starts) in enumerate(auc_prep):
                        wp = np.add.reduceat(W[:, o] * pos[o], starts, axis=1)
                        wn = np.add.reduceat(W[:, o] * neg[o], starts, axis=1)
                        below = np.cumsum(wn, axis=1) - wn
                        with np.errstate(divide="ignore", invalid="ignore"):
                            v[:, j] = (wp * (below + 0.5 * wn)).sum(1) / (wp.sum(1) * wn.sum(1))
                else:
                    continue
            draws[metric].append(v)
    rows = []
    for metric in metrics:
        if not draws[metric] or point[metric] is None:
            continue
        bs = np.vstack(draws[metric])
        for j, m in enumerate(models):
            col = bs[:, j]
            col = col[np.isfinite(col)]
            lo, hi = (np.quantile(col, [a, 1 - a]) if len(col) else (np.nan, np.nan))
            rows.append({"group": group_name, "model": m, "metric": metric, "estimate": float(point[metric][j]),
                         "ci_low": float(lo), "ci_high": float(hi), "n_eyes": int(n), "n_patients": int(K),
                         "method": method})
    return rows


def compare_models(df, task="continuous", models=None, patient=None, truth=None, settings=None, group_name="All"):
    """Compare every pair of models on the same eyes, with the patient as the unit of analysis.

    task     'continuous': model columns hold prediction errors, or predictions if truth is given;
             'binary': predicted probabilities, with truth the 0/1 diagnosis (required); 'auto': binary when the
             reference is 0/1 and every model value lies in [0, 1] (see _detect_columns)
    models, patient, truth   columns (auto-detected if None)
    settings dict overriding DEFAULT_SETTINGS

    Returns (wide, long, info). wide: one row per model pair with differences, CIs, adjusted CIs ('CI adj'),
    p-values and adjusted p-values. long: one row per pair and metric, with each model's estimate, ρΔ (rho_delta),
    the one-eye factor and the adjustments applied (p_adj_method, ci_adj_method). info: the sample, the tests, and
    each model's own performance with its patient-level CI (info["model_performance"]). Adjusted intervals are
    given only where they match the adjusted p-value (_adj_ci_matches)."""
    s = _settings(settings)
    task, models, patient, truth = _detect_columns(df, task, models, patient, truth)
    if len(models) < 2:
        raise ValueError("need at least two model columns")
    data = df.copy()
    keep = data[models].notna().all(axis=1)
    if truth is not None:
        keep &= data[truth].notna()
    if task == "continuous" and s["invalid_threshold"] is not None:
        err = data[models].values - (data[[truth]].values if truth is not None else 0)
        keep &= (np.abs(err) <= s["invalid_threshold"]).all(axis=1)
    if patient is not None:
        keep &= data[patient].notna()
    data = data[keep].reset_index(drop=True)
    if s["one_eye"] and patient is not None:
        rng_indep = np.random.default_rng(s["seed"])
        order = rng_indep.permutation(len(data))
        data = data.iloc[order].drop_duplicates(subset=[patient]).sort_index().reset_index(drop=True)
    n = len(data)
    if n < max(2, int(s["min_group_size"])):
        if s["verbose"]:
            print(f"[{group_name}] Not tested: {n} valid eye(s).")
        return pd.DataFrame(), pd.DataFrame(), {}

    use_patient = patient is not None and s["unit"] == "patient"
    cluster_ids = data[patient].astype(str).values if use_patient else np.arange(n)
    _, cluster_idx = np.unique(cluster_ids, return_inverse=True)
    K = int(cluster_idx.max()) + 1
    sizes = np.bincount(cluster_idx, minlength=K)
    has_clusters = K < n
    B = int(s["bootstrap_iterations"])
    ci_level = float(s["ci_level"]) if s["report_ci"] else None

    # patient resamples and Rademacher weights for the wild cluster bootstrap
    rng = np.random.default_rng(s["seed"])
    b_idx = np.empty((B, K), dtype=np.int32)
    for i in range(0, B, 500):
        b_idx[i:i + 500] = rng.choice(K, size=(min(500, B - i), K), replace=True)
    W_g = np.empty((B, K), dtype=np.int8)
    for i in range(0, B, 500):
        W_g[i:i + 500] = rng.choice([-1, 1], size=(min(500, B - i), K))

    pairs = list(combinations(range(len(models)), 2))
    P = data[models].values.astype(float)
    y = None if truth is None else data[truth].values.astype(float)
    metrics = s["metrics"] or (CONTINUOUS_METRICS if task == "continuous" else BINARY_METRICS)
    long_rows, rw_stats = [], {}
    # sources of the Bonferroni-type adjusted intervals: bootstrap draws (percentile CIs) and AUC standard errors
    want_adj = bool(ci_level) and s["fwer_method"] in ADJUSTED_CI_METHODS and len(pairs) > 1
    pct_draws, auc_se = {}, {}
    if s["verbose"]:
        print(f"[{group_name}] {task}: {len(models)} models, {n} eyes of {K} {'patients' if use_patient else 'units'}"
              f" ({int(np.sum(sizes >= 2))} with both eyes), {B} resamples, weighting={s['weighting']}")

    def add(metric, j, k, est_a, est_b, diff, ci, p, rho, test, factor_sizes=sizes):
        long_rows.append({"group": group_name, "metric": metric, "model_a": models[j], "model_b": models[k],
                          "est_a": est_a, "est_b": est_b, "diff": diff,
                          "ci_low": ci[0] if ci else np.nan, "ci_high": ci[1] if ci else np.nan, "p": p,
                          "rho_delta": rho,
                          "one_eye_factor": one_eye_factor(factor_sizes, rho, s["weighting"])
                          if metric in MEAN_LOSS_METRICS else one_eye_factor(factor_sizes, rho, "eye"),
                          "test": test})

    def weighted_engine(Lmat_by_metric, sqrt_keys=()):
        """Mean-loss tests with patient or minimum-variance weights, pair by pair (weights depend on rho)."""
        res = {}
        for (j, k) in pairs:
            res[(j, k)] = {}
            for f, Lm in Lmat_by_metric.items():
                rho = fellow_eye_correlation(Lm[:, j] - Lm[:, k], cluster_idx)
                w = eye_weights(cluster_idx, s["weighting"], rho)
                r, _ = all_pairs_loss_bootstrap({f: Lm[:, [j, k]]}, cluster_idx, K, b_idx, weights=w,
                                                pairs=[(0, 1)], ci_level=ci_level,
                                                sqrt_keys=tuple(x for x in sqrt_keys if x == f))
                res[(j, k)].update(r[(0, 1)])
        return res

    if task == "continuous":
        E = P if y is None else P - y[:, None]
        absE = np.abs(E)
        boot = all_pairs_bootstrap(E, cluster_idx, K, b_idx, ci_level=ci_level, prop_threshold=s["within"])
        if s["weighting"] != "eye":
            wres = weighted_engine({"MPE": E, "MAE": absE, "MSE": E ** 2}, sqrt_keys=("MSE",))
        needs_wcb = has_clusters or s["fwer_method"] == "romano_wolf"
        sd_all = wild_cluster_sd_all_pairs(E, cluster_idx, K, W_g) if needs_wcb else None
        if want_adj and {"rmse", "sd", "within"} & set(metrics):
            pct_draws = continuous_percentile_draws(E, cluster_idx, K, b_idx, s["within"], pairs)
            if s["weighting"] != "eye":                     # the RMSE interval then comes from the weighted engine
                pct_draws["rmse"] = {pr: wres[pr].get("DRAWS_MSE") for pr in pairs}
        for (j, k) in pairs:
            ea, eb = E[:, j], E[:, k]
            r = boot[(j, k)]
            fam = {"bias": ("MPE", ea, eb, ea - eb), "mae": ("MAE", np.abs(ea), np.abs(eb), np.abs(ea) - np.abs(eb)),
                   "rmse": ("RMSE", ea ** 2, eb ** 2, ea ** 2 - eb ** 2)}
            for metric, (key, la, lb, dl) in fam.items():
                if metric not in metrics:
                    continue
                rho = fellow_eye_correlation(dl, cluster_idx) if has_clusters else np.nan
                if s["weighting"] == "eye":
                    est = (np.mean(ea), np.mean(eb)) if metric == "bias" else (np.mean(la), np.mean(lb))
                    if metric == "rmse":
                        est = (np.sqrt(np.mean(la)), np.sqrt(np.mean(lb)))
                    p, ci = r[f"P_{key}"], r.get(f"CI_{key}")
                    _rw_push(rw_stats, key, r, key)
                    test = "patient bootstrap-t (CR1)" if has_clusters else "bootstrap-t"
                else:
                    wk = {"bias": "MPE", "mae": "MAE", "rmse": "MSE"}[metric]
                    wr = wres[(j, k)]
                    ea_, eb_ = wr[f"EST_{wk}"]
                    est = (np.sqrt(ea_), np.sqrt(eb_)) if metric == "rmse" else (ea_, eb_)
                    p, ci = wr[f"P_{wk}"], wr.get(f"CI_{wk}")
                    _rw_push(rw_stats, key, wr, wk)
                    test = f"patient bootstrap-t (CR1), {s['weighting']} weights"
                add(metric, j, k, est[0], est[1], est[0] - est[1], ci, p, rho, test)
            if "medae" in metrics:
                _rw_push(rw_stats, "MedAE", r, "MedAE")
                add("medae", j, k, np.median(np.abs(ea)), np.median(np.abs(eb)),
                    np.median(np.abs(ea)) - np.median(np.abs(eb)), r.get("CI_MedAE"), r["P_MedAE"], np.nan,
                    "percentile patient bootstrap" if has_clusters else "percentile bootstrap")
            if "sd" in metrics:
                if needs_wcb:
                    p_sd, obs_sd, null_sd = sd_all[(j, k)]
                    test = "Morgan-Pitman, wild cluster bootstrap"
                else:
                    p_sd, obs_sd, null_sd = comdvar(ea, eb), 0.0, np.zeros(B)
                    test = "Morgan-Pitman, HC4"
                rw_stats.setdefault("SD", {"obs": [], "null": []})
                rw_stats["SD"]["obs"].append(obs_sd); rw_stats["SD"]["null"].append(null_sd)
                add("sd", j, k, np.std(ea, ddof=1), np.std(eb, ddof=1), np.std(ea, ddof=1) - np.std(eb, ddof=1),
                    r.get("CI_SD"), p_sd, np.nan, test)
            if "within" in metrics:
                ma, mb = np.abs(ea) <= s["within"], np.abs(eb) <= s["within"]
                p_w = clustered_mcnemar_test(ma, mb, cluster_ids) if has_clusters else mcnemar_test(ma, mb)
                rho = fellow_eye_correlation(ma.astype(float) - mb.astype(float), cluster_idx) if has_clusters else np.nan
                add("within", j, k, np.mean(ma), np.mean(mb), np.mean(ma) - np.mean(mb), r.get("CI_PROP"), p_w, rho,
                    "clustered McNemar (Yang 2010)" if has_clusters else "McNemar")

    else:
        if y is None:
            raise ValueError("binary task needs a 0/1 reference standard column (truth)")
        clip = float(s["clip"])
        Pc = np.clip(P, clip, 1 - clip)
        L = {"Brier": (P - y[:, None]) ** 2,
             "LogLoss": -(y[:, None] * np.log(Pc) + (1 - y[:, None]) * np.log(1 - Pc))}
        lab = P >= s["threshold"]
        correct = lab == (y[:, None] == 1)
        subsets = {"sensitivity": y == 1, "specificity": y == 0, "accuracy": np.ones(n, bool)}
        extra = {}
        for metric, sub in subsets.items():
            extra[f"{metric}|n"] = np.bincount(cluster_idx, weights=sub.astype(float), minlength=K)
            for j in range(len(models)):
                extra[f"{metric}|{j}"] = np.bincount(cluster_idx, weights=(correct[:, j] & sub).astype(float),
                                                     minlength=K)
        loss_res, extra_sums = all_pairs_loss_bootstrap(L, cluster_idx, K, b_idx, pairs=pairs, ci_level=ci_level,
                                                        extra=extra)
        if s["weighting"] != "eye":
            wres = weighted_engine(L)
        auc = auc_all_pairs(P, y, cluster_idx, K, pairs, ci_level) if "auc" in metrics else {}
        a_ci = (1 - ci_level) / 2 if ci_level else None
        for (j, k) in pairs:
            if "auc" in metrics:
                th_j, th_k, d, lo, hi, p, se = auc[(j, k)]
                auc_se[(j, k)] = (d, se)
                add("auc", j, k, th_j, th_k, d, (lo, hi) if ci_level else None, p, np.nan,
                    "Obuchowski (1997) clustered AUC" if has_clusters else "DeLong (1988)")
            for metric, key in (("brier", "Brier"), ("logloss", "LogLoss")):
                if metric not in metrics:
                    continue
                r = loss_res[(j, k)] if s["weighting"] == "eye" else wres[(j, k)]
                rho = fellow_eye_correlation(L[key][:, j] - L[key][:, k], cluster_idx) if has_clusters else np.nan
                _rw_push(rw_stats, key, r, key)
                ea_, eb_ = r[f"EST_{key}"]
                add(metric, j, k, ea_, eb_, ea_ - eb_, r.get(f"CI_{key}"), r[f"P_{key}"], rho,
                    "patient bootstrap-t (CR1)" + ("" if s["weighting"] == "eye" else f", {s['weighting']} weights"))
            for metric, sub in subsets.items():
                if metric not in metrics:
                    continue
                ma, mb = correct[sub, j], correct[sub, k]
                ids_sub = cluster_ids[sub]
                if len(ma) < 2:
                    continue
                p_w = clustered_mcnemar_test(ma, mb, ids_sub) if has_clusters else mcnemar_test(ma, mb)
                rho = (fellow_eye_correlation(ma.astype(float) - mb.astype(float), np.unique(ids_sub, return_inverse=True)[1])
                       if has_clusters else np.nan)
                ci = None
                if ci_level:
                    with np.errstate(divide="ignore", invalid="ignore"):
                        bd = (extra_sums[f"{metric}|{j}"] - extra_sums[f"{metric}|{k}"]) / extra_sums[f"{metric}|n"]
                    bd = bd[np.isfinite(bd)]
                    ci = tuple(float(x) for x in np.quantile(bd, [a_ci, 1 - a_ci])) if len(bd) else None
                    pct_draws.setdefault(metric, {})[(j, k)] = bd
                add(metric, j, k, np.mean(ma), np.mean(mb), np.mean(ma) - np.mean(mb), ci, p_w, rho,
                    "clustered McNemar (Yang 2010)" if has_clusters else "McNemar",
                    factor_sizes=np.bincount(np.unique(ids_sub, return_inverse=True)[1]))

    long = pd.DataFrame(long_rows)
    wide = _to_wide(long, task, n, K, group_name, s)
    p_methods = {}
    wide = _adjust(wide, rw_stats, s["fwer_method"], p_methods)
    wide = _add_bonferroni_type_cis(wide, task, models, pct_draws, auc_se, s)
    wide = _drop_unmatched_adjusted_cis(wide, task, s)
    if not long.empty:
        long["p_adj"] = long.apply(lambda r: _lookup_adj(wide, r), axis=1)
        long["p_adj_method"] = [p_methods.get(WIDE_NAMES[m][1], "not available") for m in long["metric"]]
        if s["report_ci"]:
            lbl = f"{float(s['ci_level']) * 100:g}% CI"
            for side in ("low", "high"):
                long[f"ci_adj_{side}"] = long.apply(lambda r: _lookup_col(wide, r, f"{WIDE_NAMES[r['metric']][0]} {lbl} adj {side}"), axis=1)
        long["ci_adj_method"] = [_ci_adj_label(m, wide, s, has_clusters) for m in long["metric"]]
    info = {"task": task, "models": models, "patient_column": patient if use_patient else None, "truth": truth,
            "n_eyes": int(n), "n_patients": int(K), "n_bilateral": int(np.sum(sizes >= 2)),
            "bilateral_fraction": float(np.mean(sizes >= 2)), "clustered": bool(has_clusters),
            "bootstrap_iterations": B, "seed": s["seed"], "weighting": s["weighting"], "unit": s["unit"],
            "one_eye": bool(s["one_eye"]), "fwer_method": s["fwer_method"],
            "methodology": _methodology(long, task, has_clusters, s, n_pairs=len(wide)),
            "model_performance": (model_performance(P, y, task, models, metrics, cluster_idx, K, b_idx, s,
                                                    has_clusters, group_name) if s["report_ci"] else [])}
    return wide, long, info


WIDE_NAMES = {
    "bias": ("MPE Diff", "Boot (MPE) p"), "mae": ("MAE Diff", "Boot (MAE) p"),
    "medae": ("MedAE Diff", "Boot (MedAE) p"), "rmse": ("RMSE Diff", "Boot (RMSE) p"),
    "sd": ("SD Diff", "Morgan-Pitman (SD) p"), "within": ("McNemar Diff", "McNemar p"),
    "auc": ("AUC Diff", "Obuchowski (AUC) p"), "brier": ("Brier Diff", "Boot (Brier) p"),
    "logloss": ("LogLoss Diff", "Boot (LogLoss) p"), "sensitivity": ("Sensitivity Diff", "Sensitivity p"),
    "specificity": ("Specificity Diff", "Specificity p"), "accuracy": ("Accuracy Diff", "Accuracy p"),
}


def _to_wide(long, task, n, K, group_name, s):
    if long.empty:
        return pd.DataFrame()
    rows = []
    order = CONTINUOUS_METRICS if task == "continuous" else BINARY_METRICS
    for (a, b), g in long.groupby(["model_a", "model_b"], sort=False):
        row = {"Group": group_name, "Model A": a, "Model B": b, "N": n, "Patients": K}
        gm = g.set_index("metric")
        for metric in order:
            if metric not in gm.index:
                continue
            r = gm.loc[metric]
            dname, pname = WIDE_NAMES[metric]
            row[dname] = r["diff"]
            if s["report_ci"]:
                lbl = f"{float(s['ci_level']) * 100:g}% CI"
                row[f"{dname} {lbl} low"], row[f"{dname} {lbl} high"] = r["ci_low"], r["ci_high"]
            row[pname] = r["p"]
        rows.append(row)
    return pd.DataFrame(rows)


def _lookup_adj(wide, r):
    return _lookup_col(wide, r, WIDE_NAMES[r["metric"]][1].replace(" p", " p_adj"))


def _lookup_col(wide, r, col):
    sel = wide[(wide["Model A"] == r["model_a"]) & (wide["Model B"] == r["model_b"])]
    return float(sel[col].iloc[0]) if col in sel and len(sel) else np.nan


def _rw_push(rw_stats, key, r, rkey):
    """Collect one comparison's Romano-Wolf statistics: observed |t|, bootstrap null, difference and its scale."""
    st = rw_stats.setdefault(key, {"obs": [], "null": [], "diff": [], "scale": []})
    st["obs"].append(r[f"RW_{rkey}_obs"]); st["null"].append(r[f"RW_{rkey}_null"])
    st["diff"].append(r.get(f"RW_{rkey}_diff", np.nan)); st["scale"].append(r.get(f"RW_{rkey}_scale", np.nan))


METRIC_LABELS = {"bias": "Mean Error (MPE)", "mae": "MAE", "rmse": "RMSE", "medae": "MedAE", "sd": "SD",
                 "within": "Proportion Within Tolerance", "auc": "AUC", "brier": "Brier Score", "logloss": "Log-Loss",
                 "sensitivity": "Sensitivity", "specificity": "Specificity", "accuracy": "Accuracy"}


def _by_method(long, col):
    """'method: metric, metric; method: metric' from one row per metric of the long table."""
    first = long.drop_duplicates("metric")
    groups = {}
    for metric, method in zip(first["metric"], first[col]):
        groups.setdefault(method, []).append(METRIC_LABELS.get(metric, metric))
    return "; ".join(f"{k}: {', '.join(v)}" for k, v in groups.items())


def _methodology(long, task, has_clusters, s, n_pairs=None):
    if long.empty:
        return {}
    m = {"Dependence Strategy": ("Patient as the unit (cluster bootstrap / cluster-robust tests)" if has_clusters
                                 else "Independent eyes (one eye per patient or no patient identifier)"),
         "Weighting of Eyes (mean-loss metrics)": {"eye": "Every eye counts once",
                                                   "patient": "Every patient counts once",
                                                   "optimal": "Minimum-variance weights 1/(1 + (m - 1) ρΔ)"}
         [s["weighting"]]}
    for metric, g in long.groupby("metric", sort=False):
        m[f"{METRIC_LABELS.get(metric, metric)} Test"] = g["test"].iloc[0]
    if task == "binary":
        m["Classification Threshold"] = f"predicted probability >= {s['threshold']}"
    else:
        m["Tolerance (proportion within)"] = f"|error| <= {s['within']}"
    m["FWER Correction"] = ("Romano-Wolf (bootstrap tests) / Hommel (other tests)" if s["fwer_method"] == "romano_wolf"
                            else str(s["fwer_method"]))
    n_pairs = n_pairs if n_pairs is not None else long[["model_a", "model_b"]].drop_duplicates().shape[0]
    single = n_pairs == 1
    if "p_adj_method" in long:
        m["Adjusted p-values (p_adj)"] = (
            "A single model pair: no adjustment is needed, p_adj = p" if single else
            f"Adjusted for the m = {n_pairs} model pairs of each metric within the group (column p_adj_method). "
            + _by_method(long, "p_adj_method"))
    if s["report_ci"]:
        fm, level = s["fwer_method"], float(s["ci_level"])
        alpha = f"{1 - level:g}"
        if fm in ADJUSTED_CI_METHODS and single:
            text = ("A single model pair: no adjustment is needed, so the 'CI adj' columns (ci_adj_low, ci_adj_high) "
                    "equal the unadjusted intervals")
        elif fm in ("romano_wolf", "holm", "hommel"):
            boot_t = "MPE, MAE and MedAE" if task == "continuous" else "Brier score and log-loss"
            rest = ("RMSE, SD and the proportion within tolerance" if task == "continuous"
                    else "AUC, sensitivity, specificity and accuracy")
            text = ((f"Columns 'CI adj' (Estimates: ci_adj_low, ci_adj_high; method in ci_adj_method), adjusted for the "
                     f"m = {n_pairs} model pairs of each metric within the group, for the {boot_t} only: "
                     + adjusted_ci_description("romano_wolf", level) + f". The {rest} have no adjusted interval: their "
                     "adjusted p-values (Romano-Wolf from a different statistic, or Hommel) have no matching interval, "
                     "so their CIs are unadjusted")
                    if fm == "romano_wolf" else
                    f"None: {fm.capitalize()} p-values have no matching intervals, so the CIs are unadjusted")
            text += ". fwer_method='bonferroni' gives Bonferroni p-values with matching intervals for every metric"
        elif fm in ADJUSTED_CI_METHODS:
            q = f"1 - R × {alpha}/m (Benjamini-Yekutieli; R = comparisons with p_adj ≤ {alpha})" \
                if fm == "fdr_bh" else f"1 - {alpha}/m (Bonferroni)"
            boot_t, pct = (("MPE, MAE and MedAE", "RMSE, SD and the proportion within tolerance") if task == "continuous"
                           else ("Brier score and log-loss", "sensitivity, specificity and accuracy"))
            text = (f"Columns 'CI adj' (Estimates: ci_adj_low, ci_adj_high; method in ci_adj_method), adjusted for the "
                    f"m = {n_pairs} model pairs of each metric within the group. "
                    + (_by_method(long, "ci_adj_method") + ". " if "ci_adj_method" in long else "")
                    + f"{boot_t}: " + adjusted_ci_description(fm, level) + ". "
                    + (f"AUC: difference ± z × SE, z the normal quantile for coverage {q}. "
                       if task == "binary" else "")
                    + f"{pct[0].upper() + pct[1:]}: percentile interval of the same patient resamples at coverage {q}"
                    "; use more resamples (e.g. 9,999) for their far tails")
        else:
            text = adjusted_ci_description(fm, level)
        m["Adjusted Confidence Intervals"] = text
    m["Fellow-Eye Correlation"] = ("ρΔ (rho_delta): intraclass correlation of the per-eye difference in loss "
                                   "between the two models; one_eye_factor: patients a one-eye study needs, "
                                   "relative to this analysis (design effect)")
    return m


# ----------------------------------------------------------------------------------------------------------------------
# Subgroups and the full analysis
# ----------------------------------------------------------------------------------------------------------------------

QUANTILE_NAMES = {2: "half", 3: "tertile", 4: "quartile", 5: "quintile", 10: "decile"}


def _num(v):
    return f"{float(v):g}"


def _bound(edges, i):
    """Interval i of the cut points edges: '<a', 'a to <b' or '>=b'."""
    if i == 0:
        return f"<{_num(edges[0])}"
    if i == len(edges):
        return f">={_num(edges[-1])}"
    return f"{_num(edges[i - 1])} to <{_num(edges[i])}"


def define_subgroups(df, column, cuts=None, bins=None, max_levels=SUBGROUP_MAX_LEVELS):
    """Split the eyes by one column: one subgroup per level (nominal, ordinal, or numeric with at most
    max_levels values), or for a numeric column, intervals at cut points ('<22', '22 to <24.5', '>=24.5'; each
    includes its lower cut point) or equal-count quantile groups (bins; tertiles by default when there are more
    than max_levels values). Eyes with no value are in no subgroup. Returns (groups, description): groups lists
    (name, mask, definition); description has the column, its kind and a summary line."""
    if column not in df.columns:
        raise KeyError(f"subgroup column {column!r} is not in the data")
    col = df[column]
    present = np.array(col.notna(), dtype=bool)
    is_cat = isinstance(col.dtype, pd.CategoricalDtype)
    numeric = pd.api.types.is_numeric_dtype(col) and not pd.api.types.is_bool_dtype(col) and not is_cat
    if (cuts is not None or bins is not None) and not numeric:
        x = pd.to_numeric(col.astype(object), errors="coerce")
        if x.notna().sum() < present.sum():
            raise ValueError(f"subgroup_cuts and subgroup_bins need a numeric subgroup column; {column!r} is not")
        col, numeric = x, True
    groups, default_bins = [], False
    if numeric:
        x = col.astype("float64").to_numpy()
        present = present & np.isfinite(x)
        levels = np.unique(x[present])
        if cuts is None and bins is None and len(levels) > max_levels:
            bins, default_bins = 3, True
        if cuts is None and bins is None:
            kind = "ordinal (numeric levels)"
            for v in levels:
                name = f"{column} = {_num(v)}"
                groups.append((name, present & (x == v), name))
            summary = f"{column}, {kind}: one subgroup per value ({', '.join(_num(v) for v in levels)})"
        elif cuts is not None:
            edges = sorted({float(c) for c in np.atleast_1d(cuts)})
            kind = "continuous, cut points"
            idx = np.searchsorted(edges, x, side="right")
            for i in range(len(edges) + 1):
                name = f"{column} {_bound(edges, i)}"
                groups.append((name, present & (idx == i), name))
            summary = (f"{column}, continuous, cut points {', '.join(_num(c) for c in edges)}: "
                       f"{', '.join(_bound(edges, i) for i in range(len(edges) + 1))} (each interval includes its "
                       "lower cut point)")
        else:
            q = int(bins)
            edges = list(np.unique(np.quantile(x[present], np.arange(1, q) / q))) if present.any() else []
            idx = np.searchsorted(edges, x, side="right")
            filled = [i for i in range(len(edges) + 1) if np.any(present & (idx == i))]
            word = QUANTILE_NAMES.get(q) if len(filled) == q else None
            kind = "continuous, equal-count quantile groups"
            for r, i in enumerate(filled, 1):
                mask = present & (idx == i)
                name = (f"{column} {word or 'quantile group'} {r} of {len(filled)} "
                        f"({_num(x[mask].min())} to {_num(x[mask].max())})")
                groups.append((name, mask, f"{name}, i.e. {column} {_bound(edges, i)}"))
            summary = (f"{column}, continuous: {len(filled)} equal-count quantile groups"
                       + (f" ({word}s)" if word else "")
                       + (f" cut at {', '.join(_num(e) for e in edges)}" if edges else "")
                       + (f"; tertiles by default, as the column has more than {max_levels} distinct values "
                          "(set subgroup_cuts or subgroup_bins, CLI --cuts or --bins, to choose)" if default_bins else "")
                       + ("" if len(filled) == q else f"; fewer than the {q} groups asked for, because of tied values"))
    else:
        if is_cat:
            levels = [c for c in col.cat.categories if (col == c).any()]
            kind = "ordinal (ordered categories)" if col.cat.ordered else "nominal"
        else:
            levels = sorted(col.dropna().unique(), key=str)
            kind = "nominal"
        for lvl in levels:
            name = f"{column} = {lvl}"
            groups.append((name, present & (col == lvl).values, name))
        summary = f"{column}, {kind}: one subgroup per level ({', '.join(str(v) for v in levels)})"
    n_missing = int(len(df) - present.sum())
    if n_missing:
        summary += f"; {n_missing} eye(s) with no {column} value are in All only"
    groups = [g for g in groups if g[1].any()]
    return groups, {"column": column, "kind": kind, "summary": summary, "default_bins": default_bins,
                    "n_missing": n_missing}


_SPEC_RE = re.compile(r"^(?P<col>.+?):(?P<key>bins|cuts)=(?P<val>[^:]+)$")


def parse_subgroup_spec(spec):
    """One subgroup definition as {'column', 'cuts', 'bins'}. Accepted forms:
      'eye'                           a column, split as define_subgroups does by default
      'axial_length:cuts=22,24.5,26'  cut points (each interval includes its lower cut point)
      'mean_k:bins=5'                 equal-count quantile groups (here quintiles)
      ('mean_k', {'bins': 5})  or  {'column': 'mean_k', 'bins': 5}
    """
    if isinstance(spec, dict):
        opts = dict(spec)
        col = opts.pop("column", None)
    elif isinstance(spec, tuple) and len(spec) == 2 and isinstance(spec[1], dict):
        col, opts = spec[0], dict(spec[1])
    elif isinstance(spec, str):
        m = _SPEC_RE.match(spec.strip())
        if m:
            col, key, val = m["col"], m["key"], m["val"]
            try:
                opts = {key: [float(v) for v in val.split(",") if v.strip()] if key == "cuts" else int(val)}
            except ValueError:
                raise ValueError(f"cannot read {key!r} in subgroup {spec!r}: use e.g. 'axial_length:cuts=22,24.5' "
                                 "or 'mean_k:bins=5'") from None
        else:
            col, opts = spec, {}
    else:
        raise TypeError(f"a subgroup is a column name, 'column:bins=Q', 'column:cuts=a,b', (column, {{...}}) or "
                        f"{{'column': ...}}; got {spec!r}")
    if col is None or col == "":
        raise ValueError(f"subgroup {spec!r} names no column")
    unknown = set(opts) - {"cuts", "bins"}
    if unknown:
        raise ValueError(f"subgroup {col!r}: unknown option(s) {sorted(unknown)}; use 'cuts' or 'bins'")
    cuts, bins = opts.get("cuts"), opts.get("bins")
    if cuts is not None and bins is not None:
        raise ValueError(f"subgroup {col!r}: give cuts or bins, not both")
    if cuts is not None:
        cuts = sorted({float(c) for c in np.atleast_1d(cuts)})
        if not cuts or not np.all(np.isfinite(cuts)):
            raise ValueError(f"subgroup {col!r}: cuts must be one or more finite numbers")
    if bins is not None:
        if int(bins) != bins or int(bins) < 2:
            raise ValueError(f"subgroup {col!r}: bins must be an integer of at least 2")
        bins = int(bins)
    return {"column": col, "cuts": cuts, "bins": bins}


def _is_single_spec(x):
    return isinstance(x, (str, dict)) or (isinstance(x, tuple) and len(x) == 2 and isinstance(x[1], dict))


def _subgroup_specs(subgroup, s):
    """The list of subgroup definitions; settings subgroup_cuts/subgroup_bins apply to a single column name."""
    if subgroup is None:
        return []
    specs = [parse_subgroup_spec(x) for x in ([subgroup] if _is_single_spec(subgroup) else list(subgroup))]
    if s["subgroup_cuts"] is not None or s["subgroup_bins"] is not None:
        if len(specs) != 1 or specs[0]["cuts"] is not None or specs[0]["bins"] is not None:
            raise ValueError("settings subgroup_cuts and subgroup_bins apply to one subgroup given by its column name; "
                             "with several subgroups give the cut points or bins in each, e.g. "
                             "'axial_length:cuts=22,24.5' or 'mean_k:bins=5'")
        specs[0]["cuts"], specs[0]["bins"] = s["subgroup_cuts"], s["subgroup_bins"]
    return specs


def _blank_tests(w, l, info, why):
    """Keep the estimates of a group but remove its intervals and tests (too few patients)."""
    w = w.copy()
    cols = [c for c in w.columns if c.endswith(" p") or c.endswith(" p_adj") or "95% CI" in c]
    w[cols] = np.nan
    l = l.copy()
    for c in ("ci_low", "ci_high", "p", "p_adj", "ci_adj_low", "ci_adj_high"):
        if c in l.columns:
            l[c] = np.nan
    l["test"] = why
    for c in ("p_adj_method", "ci_adj_method"):
        if c in l.columns:
            l[c] = ""
    info = dict(info)
    info["model_performance"] = [{**r, "ci_low": np.nan, "ci_high": np.nan, "method": "point estimate only (" + why + ")"}
                                 for r in info.get("model_performance", [])]
    return w, l, info


def analyze(df, task="auto", models=None, patient=None, truth=None, subgroup=None, cross=None, settings=None):
    """compare_models for all eyes and for each subgroup; returns the wide table with the long table, each
    model's performance, the methodology and the sample in .attrs, for write_results_to_excel.

    subgroup  a definition or a list of them, each analyzed separately: a column of any type ('eye'), or a numeric
              column with cut points or equal-count groups ('axial_length:cuts=22,24.5', 'mean_k:bins=5'; see
              parse_subgroup_spec). For a single column name, settings subgroup_cuts or subgroup_bins also work.
    cross     optional [rows, columns] definitions crossed with each other (off by default); cells with fewer than
              settings['cross_min_patients'] patients get point estimates only. Adds the 'Grid' sheet.

    Models and the patient and reference columns are found once, on all eyes. Each group is adjusted for its own
    model pairs, so subgroup results are exploratory."""
    s = _settings(settings)
    specs = _subgroup_specs(subgroup, s)
    cross_specs = None
    if cross is not None:
        cross_specs = [parse_subgroup_spec(x) for x in ([cross] if _is_single_spec(cross) else list(cross))]
        if len(cross_specs) != 2:
            raise ValueError("cross takes two subgroup definitions: [rows, columns]")
        if cross_specs[0]["column"] == cross_specs[1]["column"]:
            raise ValueError("cross needs two different columns")
    sub_cols = list(dict.fromkeys([d["column"] for d in specs] + [d["column"] for d in (cross_specs or [])]))
    for c in sub_cols:
        if models is not None and c in list(models):
            raise ValueError(f"the subgroup column {c!r} cannot also be a model")
    task, models, patient, truth = _detect_columns(df, task, models, patient, truth, extra_ignore=tuple(sub_cols))
    groups, descs, seen = [("All", np.ones(len(df), bool), "All eyes", {"grouping": "All"})], [], {"All"}

    def add(name, mask, definition, meta):
        if name not in seen:
            seen.add(name)
            groups.append((name, mask, definition, meta))

    for spec in specs:
        g, d = define_subgroups(df, spec["column"], spec["cuts"], spec["bins"])
        descs.append(d)
        if s["verbose"] and d["default_bins"]:
            print(f"[subgroups] {spec['column']} has more than {SUBGROUP_MAX_LEVELS} distinct values: split into "
                  "tertiles (equal numbers of eyes) by default; give cut points or bins to choose (e.g. "
                  f"'{spec['column']}:bins=5').")
        for name, mask, definition in g:
            add(name, mask, definition, {"grouping": spec["column"]})
    grid, cross_desc = None, None
    if cross_specs:
        (ga, da), (gb, db) = [define_subgroups(df, c["column"], c["cuts"], c["bins"]) for c in cross_specs]
        label = f"{da['column']} × {db['column']}"
        cross_desc = (da, db)
        grid = {"label": label, "row_column": da["column"], "col_column": db["column"],
                "rows": [n for n, _, _ in ga], "cols": [n for n, _, _ in gb], "cells": {},
                "min_patients": int(s["cross_min_patients"])}
        for na, ma, _ in ga:
            for nb, mb, _ in gb:
                name = f"{na} × {nb}"
                add(name, ma & mb, name, {"grouping": label, "row": na, "col": nb, "crossed": True})
    wides, longs, infos, sample, gmap = [], [], {}, {}, {}
    for name, mask, definition, meta in groups:
        gmap[name] = meta["grouping"]
        sub = df[mask]
        n_pat = int(sub[patient].nunique()) if patient is not None else len(sub)
        if meta.get("crossed"):
            grid["cells"][name] = {"row": meta["row"], "col": meta["col"], "n_eyes": int(len(sub)),
                                   "n_patients": n_pat, "tested": False}
        if len(sub) == 0:
            sample[name] = {"grouping": meta["grouping"], "definition": f"{definition}: no eyes"}
            continue
        w, l, info = compare_models(sub, task=task, models=models, patient=patient, truth=truth, settings=s,
                                    group_name=name)
        if w.empty:
            sample[name] = {"grouping": meta["grouping"],
                            "definition": f"{definition}: {len(sub)} eye(s), not tested (fewer than "
                                          f"{max(2, int(s['min_group_size']))} valid eyes)"}
            continue
        note = ""
        if meta.get("crossed"):
            if n_pat < int(s["cross_min_patients"]):
                why = f"fewer than {int(s['cross_min_patients'])} patients"
                w, l, info = _blank_tests(w, l, info, why)
                note = f"; point estimates only ({why})"
            else:
                grid["cells"][name]["tested"] = True
        units = (f"{info['n_patients']} patients" if info["patient_column"] is not None
                 else "eyes analyzed as independent units")
        wides.append(w); longs.append(l)
        infos[name] = info
        sample[name] = {"grouping": meta["grouping"], "definition": f"{definition}: {info['n_eyes']} eyes, {units}{note}",
                        **{k: v for k, v in info.items() if k not in ("methodology", "model_performance")}}
    if not wides:
        return pd.DataFrame()
    final = pd.concat(wides, ignore_index=True)
    long_all = pd.concat(longs, ignore_index=True)
    long_all.insert(1, "grouping", long_all["group"].map(gmap))
    first = next(iter(infos.values()))
    methodology = dict(first["methodology"])
    if descs:
        methodology["Subgroups"] = ("; ".join(d["summary"] for d in descs) + ". Each subgroup is analyzed separately "
                                    "with the same models, tests and settings")
    if cross_desc:
        da, db = cross_desc
        methodology["Crossed subgroups"] = (f"{da['summary']}, crossed with {db['summary']}. Cut points come from all "
                                            "eyes, so a row or column covers the same range in every cell. Cells with "
                                            f"fewer than {int(s['cross_min_patients'])} patients get point estimates "
                                            "only, without intervals or tests")
    if descs or cross_desc:
        methodology["Subgroup adjustment"] = ("Subgroup results are exploratory: p-values and intervals are adjusted "
                                              "for the model pairs within each group, not across groups")
    final.attrs["methodology_summary"] = [{"Test Category": k, "Methodology Applied": v} for k, v in methodology.items()]
    final.attrs["long"] = long_all.to_dict("records")
    mp = [{**r, "grouping": gmap.get(r.get("group"), "")} for v in infos.values() for r in v.get("model_performance", [])]
    final.attrs["model_performance"] = [{k: r[k] for k in ["group", "grouping"] + [c for c in r if c not in ("group", "grouping")]}
                                        for r in mp]
    final.attrs["group_info"] = sample
    if grid is not None:
        grid["task"] = task
        final.attrs["grid"] = grid
    return final


# ----------------------------------------------------------------------------------------------------------------------
# Excel output
# ----------------------------------------------------------------------------------------------------------------------

GRID_METRIC_LABELS = {"bias": "Mean error", "mae": "Mean absolute error", "rmse": "Root mean squared error",
                      "medae": "Median absolute error", "sd": "SD of the error", "within": "Proportion within tolerance",
                      "auc": "AUC", "brier": "Brier score", "logloss": "Log-loss", "sensitivity": "Sensitivity",
                      "specificity": "Specificity", "accuracy": "Accuracy"}
LOWER_IS_BETTER = {"mae", "rmse", "medae", "sd", "brier", "logloss"}
HIGHER_IS_BETTER = {"within", "auc", "sensitivity", "specificity", "accuracy"}


def _set_column_widths(writer, sheet_name, widths):
    """Column widths for an openpyxl or an xlsxwriter ExcelWriter sheet."""
    try:
        ws = writer.sheets[sheet_name]
        if hasattr(ws, 'set_column'):                                   # xlsxwriter
            for idx, w in enumerate(widths):
                ws.set_column(idx, idx, w)
        else:                                                           # openpyxl
            from openpyxl.utils import get_column_letter
            for idx, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(idx)].width = w
    except (ImportError, AttributeError, KeyError):
        pass


def _write_results_sheet(results, writer):
    """The comparisons ('Results') and the tests and adjustments used ('Statistical Tests')."""
    if results.empty:
        return
    results.to_excel(writer, sheet_name="Results", index=False)
    _set_column_widths(writer, "Results", [15 if " p" in c else 22 for c in results.columns])
    summary = results.attrs.get("methodology_summary")
    if summary:
        meth = pd.DataFrame(summary)
        meth.to_excel(writer, sheet_name="Statistical Tests", index=False)
        _set_column_widths(writer, "Statistical Tests", [35] + [75] * (len(meth.columns) - 1))


def _write_grid_sheet(path, results):
    """Grid sheet for crossed subgroups: one block per metric and model, with a color scale shared by the
    models of a metric (white best, red worst; for the mean error, red furthest from 0)."""
    grid = results.attrs.get("grid")
    mp = pd.DataFrame(results.attrs.get("model_performance", []))
    if not grid or mp.empty:
        return
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    cells = grid["cells"]
    mp = mp[mp["group"].isin(list(cells))]
    if mp.empty:
        return
    est = {(r.group, r.model, r.metric): r.estimate for r in mp.itertuples()}
    models = list(dict.fromkeys(mp["model"]))
    metrics = list(dict.fromkeys(mp["metric"]))
    rows, cols = grid["rows"], grid["cols"]
    by_rc = {(c["row"], c["col"]): name for name, c in cells.items()}
    wb = load_workbook(path)
    ws = wb.create_sheet("Grid")
    bold, grey = Font(bold=True), Font(italic=True, color="808080")
    ws.cell(1, 1, f"{grid['label']}: each model's performance in each cell (rows: {grid['row_column']}; columns: "
                  f"{grid['col_column']}). Cut points come from all eyes. Gray italics: fewer than "
                  f"{grid['min_patients']} patients (point estimate only, not colored). Empty: no eyes. Within each metric the "
                  "color scale is shared by all models: white is best and red worst (for the mean error, red is "
                  "furthest from 0).").font = Font(italic=True)
    width = len(cols) + 2                                          # row labels, columns, one blank column

    def block(r0, c0, title, value_of, fmt):
        ws.cell(r0, c0, title).font = bold
        for j, cl in enumerate(cols):
            h = ws.cell(r0 + 1, c0 + 1 + j, cl)
            h.font = bold
            h.alignment = Alignment(wrap_text=True)
        refs = []
        for i, rw in enumerate(rows):
            ws.cell(r0 + 2 + i, c0, rw).font = bold
            for j, cl in enumerate(cols):
                name = by_rc.get((rw, cl))
                v = value_of(name) if name is not None else None
                if v is None or (isinstance(v, float) and not np.isfinite(v)):
                    continue
                cell = ws.cell(r0 + 2 + i, c0 + 1 + j, v)
                if fmt:
                    cell.number_format = fmt
                if name is not None and not cells[name]["tested"]:
                    cell.font = grey                                   # shown, but kept out of the color scale
                else:
                    refs.append(cell.coordinate)
        return refs

    r = 3
    block(r, 1, "Patients (eyes) per cell", lambda n: (f"{cells[n]['n_patients']} ({cells[n]['n_eyes']})"
                                                         if cells[n]["n_eyes"] else None), None)
    r += len(rows) + 3
    for metric in metrics:
        sense = ("lower is better" if metric in LOWER_IS_BETTER else "higher is better" if metric in HIGHER_IS_BETTER
                 else "closer to 0 is better")
        ws.cell(r, 1, f"{GRID_METRIC_LABELS.get(metric, metric)} ({sense})").font = Font(bold=True, size=12)
        refs = []
        for k, model in enumerate(models):
            refs += block(r + 1, 1 + k * width, str(model), lambda n, m=model: est.get((n, m, metric)), "0.000")
        vals = [v for (g, m, mt), v in est.items() if mt == metric and np.isfinite(v) and cells[g]["tested"]]
        if refs and vals:
            rng_ = " ".join(refs)
            if metric in LOWER_IS_BETTER:
                rule = ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="D6604D")
            elif metric in HIGHER_IS_BETTER:
                rule = ColorScaleRule(start_type="min", start_color="D6604D", end_type="max", end_color="FFFFFF")
            else:
                big = max(abs(v) for v in vals) or 1.0
                rule = ColorScaleRule(start_type="num", start_value=-big, start_color="D6604D", mid_type="num",
                                      mid_value=0, mid_color="FFFFFF", end_type="num", end_value=big,
                                      end_color="D6604D")
            ws.conditional_formatting.add(rng_, rule)
        r += len(rows) + 4
    ws.column_dimensions["A"].width = 30
    for k in range(len(models)):
        c0 = 1 + k * width
        ws.column_dimensions[get_column_letter(c0)].width = 30
        for j in range(len(cols)):
            ws.column_dimensions[get_column_letter(c0 + 1 + j)].width = 16
    wb.save(path)


def write_results_to_excel(results, path):
    """Write the sheets Results, Statistical Tests, Model Performance (each model's performance with
    its patient-level CI), Estimates (every comparison with ρΔ, the one-eye factor and the adjustments) and
    Sample, plus Grid with crossed subgroups."""
    with pd.ExcelWriter(path) as writer:
        _write_results_sheet(results, writer)
        if results.attrs.get("model_performance"):
            mp = pd.DataFrame(results.attrs["model_performance"])
            mp.to_excel(writer, sheet_name="Model Performance", index=False)
            _set_column_widths(writer, "Model Performance", [18] * len(mp.columns))
        if "long" in results.attrs:
            est = pd.DataFrame(results.attrs["long"])
            est.to_excel(writer, sheet_name="Estimates", index=False)
            _set_column_widths(writer, "Estimates", [36 if c in ("test", "p_adj_method", "ci_adj_method") else 18
                                                     for c in est.columns])
        if "group_info" in results.attrs:
            gi = pd.DataFrame(results.attrs["group_info"]).T.reset_index().rename(columns={"index": "Group"})
            gi.to_excel(writer, sheet_name="Sample", index=False)
            _set_column_widths(writer, "Sample", [30, 24, 60] + [16] * (len(gi.columns) - 3))
    if results.attrs.get("grid"):
        _write_grid_sheet(path, results)


# ----------------------------------------------------------------------------------------------------------------------
# Command line
# ----------------------------------------------------------------------------------------------------------------------

def _command_line(argv=None):
    import argparse
    ap = argparse.ArgumentParser(
        description="Compare prediction models evaluated on eyes (one or both eyes per patient), with the patient as "
                    "the unit of analysis. Continuous outputs: give one column of prediction errors (prediction minus "
                    "reference) per model, or predictions with --truth. Classifiers: one column of predicted "
                    "probabilities per model and the 0/1 diagnosis (--truth), which AUC, sensitivity and specificity "
                    "need.")
    ap.add_argument("csv", help="one row per eye: a patient ID column and one column of prediction errors per model "
                                "(continuous outputs); or predictions plus a reference column (--truth); or, for "
                                "classifiers, predicted probabilities plus the 0/1 diagnosis (--truth)")
    ap.add_argument("--task", default="auto", choices=["auto", "continuous", "binary"],
                    help="auto (default): binary when the reference is 0/1 and every model value lies in [0, 1], "
                         "else continuous (the model columns are prediction errors when there is no reference)")
    ap.add_argument("--models", nargs="*", default=None,
                    help="model columns (default: every numeric column other than the patient, reference, subgroup "
                         "and common metadata columns)")
    ap.add_argument("--patient", default=None, help="patient ID column (default: patient_id, id, subject, ...)")
    ap.add_argument("--truth", default=None,
                    help="reference column: optional for continuous outputs (the model columns are then predictions "
                         "and the errors prediction - truth), required for classifiers (the 0/1 diagnosis). With "
                         "--task auto or binary a column named truth, reference, diagnosis, measured, ... is used "
                         "if present")
    ap.add_argument("--subgroup", action="append", default=None, metavar="SPEC",
                    help="also give results for each subgroup of a column; repeat for several columns, each analyzed "
                         "separately. SPEC is a column (nominal, e.g. eye; ordinal; or continuous, split into "
                         "tertiles by default) or COLUMN:cuts=A,B,... (cut points) or COLUMN:bins=Q (Q equal-count "
                         "groups), e.g. --subgroup axial_length:cuts=22,24.5,26 --subgroup mean_k:bins=5")
    split = ap.add_mutually_exclusive_group()
    split.add_argument("--cuts", nargs="+", type=float, default=None, metavar="CUT",
                       help="with a single --subgroup COLUMN: cut points, e.g. --cuts 22 24.5 gives <22, 22 to <24.5 "
                            "and >=24.5")
    split.add_argument("--bins", type=int, default=None, metavar="Q",
                       help="with a single --subgroup COLUMN: Q equal-count quantile groups (e.g. 3 for tertiles)")
    ap.add_argument("--cross", nargs=2, default=None, metavar=("ROWS", "COLUMNS"),
                    help="optional: cross two subgroup SPECs, e.g. --cross axial_length:cuts=22,26 mean_k:bins=3 "
                         "(3 x 3 cells), with a Grid sheet for heat maps. Cut points come from all eyes")
    ap.add_argument("--cross-min-patients", type=int, default=20, metavar="N",
                    help="crossed cells with fewer patients get point estimates only (default 20)")
    ap.add_argument("--weighting", default="eye", choices=["eye", "patient", "optimal"])
    ap.add_argument("--threshold", type=float, default=0.5, help="classifiers: probability threshold (default 0.5)")
    ap.add_argument("--within", type=float, default=MCNEMAR_THRESHOLD,
                    help="continuous outputs: tolerance for the proportion of eyes with |error| <= within")
    ap.add_argument("--out", default="ophthalmology_stats_results.xlsx")
    a = ap.parse_args(argv)
    if (a.cuts is not None or a.bins is not None) and (a.subgroup is None or len(a.subgroup) != 1):
        ap.error("--cuts and --bins go with a single --subgroup COLUMN; for several subgroups write the cut points or "
                 "bins in each, e.g. --subgroup axial_length:cuts=22,24.5 --subgroup mean_k:bins=5")
    sub = None if a.subgroup is None else (a.subgroup[0] if len(a.subgroup) == 1 else a.subgroup)
    res = analyze(pd.read_csv(a.csv), task=a.task, models=a.models, patient=a.patient, truth=a.truth,
                  subgroup=sub, cross=a.cross, settings={"weighting": a.weighting, "threshold": a.threshold,
                                                         "within": a.within, "subgroup_cuts": a.cuts,
                                                         "subgroup_bins": a.bins,
                                                         "cross_min_patients": a.cross_min_patients})
    write_results_to_excel(res, a.out)
    print(f"Results written to {a.out}")
    return res


if __name__ == "__main__":
    _command_line()
